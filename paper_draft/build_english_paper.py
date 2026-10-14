from __future__ import annotations

import json
import math
from pathlib import Path

import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = Path(__file__).resolve().parent
ASSET_DIR = OUT_DIR / "assets"
OUTPUT = OUT_DIR / "comsol_trustworthy_surrogate_paper_draft_en.docx"

NAVY = "17365D"
MID_BLUE = "4F81BD"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9E1F2"
BORDER = "D9D9D9"
TEXT = "111111"


def font(path: str, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(path, size=size)


ARIAL = "C:/Windows/Fonts/arial.ttf"
ARIAL_BOLD = "C:/Windows/Fonts/arialbd.ttf"


def draw_centered(draw: ImageDraw.ImageDraw, xy: tuple[float, float], text: str,
                  fnt: ImageFont.FreeTypeFont, fill: str = "#111111") -> None:
    box = draw.multiline_textbbox((0, 0), text, font=fnt, align="center", spacing=4)
    w = box[2] - box[0]
    h = box[3] - box[1]
    draw.multiline_text((xy[0] - w / 2, xy[1] - h / 2), text,
                        font=fnt, fill=fill, align="center", spacing=4)


def draw_vertical_label(img: Image.Image, xy: tuple[float, float], text: str,
                        fnt: ImageFont.FreeTypeFont, fill: str = "#222222") -> None:
    """Draw a centered, counter-clockwise axis label without clipping it."""
    probe = ImageDraw.Draw(img)
    box = probe.textbbox((0, 0), text, font=fnt)
    width = box[2] - box[0] + 24
    height = box[3] - box[1] + 24
    label = Image.new("RGBA", (width, height), (255, 255, 255, 0))
    label_draw = ImageDraw.Draw(label)
    label_draw.text((12 - box[0], 12 - box[1]), text, font=fnt, fill=fill)
    rotated = label.rotate(90, expand=True, resample=Image.Resampling.BICUBIC)
    img.paste(rotated,
              (round(xy[0] - rotated.width / 2), round(xy[1] - rotated.height / 2)),
              rotated)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int],
          color: str = "#4F81BD", width: int = 5) -> None:
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 15
    spread = 0.55
    p1 = (end[0] - length * math.cos(angle - spread), end[1] - length * math.sin(angle - spread))
    p2 = (end[0] - length * math.cos(angle + spread), end[1] - length * math.sin(angle + spread))
    draw.polygon([end, p1, p2], fill=color)


def make_workflow_figure(path: Path) -> None:
    img = Image.new("RGB", (1800, 900), "white")
    draw = ImageDraw.Draw(img)
    title_font = font(ARIAL_BOLD, 40)
    box_font = font(ARIAL_BOLD, 25)
    sub_font = font(ARIAL, 20)
    draw_centered(draw, (900, 52), "Traceable requirement driven automatic modeling workflow", title_font)

    boxes = [
        ("1  Requirement contract", "Geometry materials boundaries\nstudy outputs and units"),
        ("2  Safety gate", "Missing decisive inputs and\nexplicit conflicts block execution"),
        ("3  Physics selection", "Interfaces coupling and\nstudy are selected hierarchically"),
        ("4  Model plan", "Parameters geometry materials\nboundaries mesh and solver"),
        ("5  Code and static audit", "Mapped MATLAB builder plus\nparameter and selection checks"),
        ("6  Approved execution", "MATLAB LiveLink builds meshes\nsolves and exports evidence"),
        ("7  Physics validation", "Bounds conservation and\ncase specific invariants"),
        ("8  Evidence package", "MPH CSV preview stage log and\ndownstream design or surrogate use"),
    ]
    bw, bh = 380, 155
    x_positions = [35, 475, 915, 1355]
    positions = [
        (x_positions[0], 115),
        (x_positions[1], 115),
        (x_positions[2], 115),
        (x_positions[3], 115),
        (x_positions[3], 365),
        (x_positions[2], 365),
        (x_positions[1], 365),
        (x_positions[0], 365),
    ]
    for (heading, detail), (x, y) in zip(boxes, positions):
        draw.rounded_rectangle((x, y, x + bw, y + bh), radius=18,
                               fill="#EAF2F8", outline="#4F81BD", width=4)
        draw_centered(draw, (x + bw / 2, y + 43), heading, box_font, "#17365D")
        draw_centered(draw, (x + bw / 2, y + 108), detail, sub_font, "#333333")

    arrow(draw, (420, 192), (465, 192))
    arrow(draw, (860, 192), (905, 192))
    arrow(draw, (1300, 192), (1345, 192))
    arrow(draw, (1545, 275), (1545, 355))
    arrow(draw, (1345, 442), (1300, 442))
    arrow(draw, (905, 442), (860, 442))
    arrow(draw, (465, 442), (420, 442))

    draw.rounded_rectangle((300, 645, 1500, 825), radius=22,
                           fill="#F6F6F6", outline="#7F8C8D", width=4)
    draw_centered(draw, (900, 694), "Controlled failure and repair path", box_font, "#333333")
    draw_centered(draw, (900, 765),
                  "Return unanswered requirements for clarification; retry only an allowlisted failure once;\notherwise preserve the failed artifact and require engineering review",
                  sub_font, "#333333")
    draw.line([(225, 525), (225, 595), (900, 595)], fill="#7F8C8D", width=5)
    arrow(draw, (900, 595), (900, 635), "#7F8C8D", 5)
    img.save(path, dpi=(220, 220))


def interpolate_color(value: float) -> tuple[int, int, int]:
    stops = [
        (0.0, (32, 92, 159)),
        (0.25, (93, 165, 219)),
        (0.5, (247, 247, 247)),
        (0.75, (244, 151, 94)),
        (1.0, (178, 24, 43)),
    ]
    value = max(0.0, min(1.0, value))
    for (p0, c0), (p1, c1) in zip(stops[:-1], stops[1:]):
        if p0 <= value <= p1:
            t = (value - p0) / (p1 - p0)
            return tuple(round(c0[j] + t * (c1[j] - c0[j])) for j in range(3))
    return stops[-1][1]


def make_response_heatmap(csv_path: Path, path: Path) -> None:
    df = pd.read_csv(csv_path)
    pivot = df.pivot(index="htc_W_m2K", columns="Vtot_mV", values="Tmax_C").sort_index(ascending=False)
    values = pivot.to_numpy()
    vmin, vmax = float(values.min()), float(values.max())
    img = Image.new("RGB", (1500, 1050), "white")
    draw = ImageDraw.Draw(img)
    title_font = font(ARIAL_BOLD, 40)
    axis_font = font(ARIAL_BOLD, 28)
    tick_font = font(ARIAL, 23)
    cell_font = font(ARIAL, 19)
    draw_centered(draw, (750, 55), "Maximum temperature across the 49 point COMSOL sweep", title_font)

    left, top, cell_w, cell_h = 185, 130, 140, 105
    rows, cols = values.shape
    for i in range(rows):
        for j in range(cols):
            val = float(values[i, j])
            frac = (val - vmin) / (vmax - vmin)
            color = interpolate_color(frac)
            x1, y1 = left + j * cell_w, top + i * cell_h
            draw.rectangle((x1, y1, x1 + cell_w, y1 + cell_h), fill=color,
                           outline="#FFFFFF", width=3)
            text_color = "#FFFFFF" if frac < 0.24 or frac > 0.76 else "#111111"
            draw_centered(draw, (x1 + cell_w / 2, y1 + cell_h / 2), f"{val:.1f}", cell_font, text_color)

    for j, value in enumerate(pivot.columns):
        draw_centered(draw, (left + j * cell_w + cell_w / 2, top + rows * cell_h + 35),
                      f"{float(value):g}", tick_font)
    for i, value in enumerate(pivot.index):
        draw.text((left - 25, top + i * cell_h + cell_h / 2), f"{float(value):g}",
                  font=tick_font, fill="#222222", anchor="rm")
    draw_centered(draw, (left + cols * cell_w / 2, 945), "Applied voltage  Vtot  mV", axis_font)
    draw_vertical_label(img, (55, top + rows * cell_h / 2),
                        "Convection coefficient  h  W m-2 K-1", axis_font)
    legend_x = 1230
    legend_top = top
    legend_h = rows * cell_h
    for y in range(legend_h):
        frac = 1 - y / max(1, legend_h - 1)
        draw.line((legend_x, legend_top + y, legend_x + 55, legend_top + y),
                  fill=interpolate_color(frac), width=1)
    draw.rectangle((legend_x, legend_top, legend_x + 55, legend_top + legend_h),
                   outline="#777777", width=2)
    draw.text((legend_x + 70, legend_top), f"{vmax:.1f}", font=tick_font, fill="#222222", anchor="lm")
    draw.text((legend_x + 70, legend_top + legend_h), f"{vmin:.1f}", font=tick_font, fill="#222222", anchor="lm")
    draw.text((legend_x - 5, legend_top - 38), "Tmax  degC", font=axis_font, fill="#222222")
    img.save(path, dpi=(220, 220))


def make_model_comparison(record_path: Path, path: Path) -> None:
    record = json.loads(record_path.read_text(encoding="utf-8"))
    reports = record["report"]["candidate_reports"]
    labels = ["Mean\nbaseline", "Ridge", "Random\nForest", "Scaled\nMLP"]
    rmse = [float(item["test_rmse"]) for item in reports]
    r2 = [float(item["test_r2"]) for item in reports]
    img = Image.new("RGB", (1450, 900), "white")
    draw = ImageDraw.Draw(img)
    title_font = font(ARIAL_BOLD, 40)
    axis_font = font(ARIAL_BOLD, 27)
    tick_font = font(ARIAL, 22)
    label_font = font(ARIAL_BOLD, 24)
    draw_centered(draw, (725, 55), "Candidate surrogate performance on the internal test split", title_font)
    left, top, right, bottom = 155, 135, 1330, 725
    draw.line((left, top, left, bottom), fill="#555555", width=3)
    draw.line((left, bottom, right, bottom), fill="#555555", width=3)
    ymax = 30
    for yval in range(0, 31, 5):
        y = bottom - (yval / ymax) * (bottom - top)
        draw.line((left, y, right, y), fill="#E1E1E1", width=2)
        draw.text((left - 18, y), str(yval), font=tick_font, fill="#333333", anchor="rm")
    colors = ["#95A5A6", "#7FB3D5", "#4F81BD", "#17365D"]
    slot = (right - left) / len(rmse)
    for i, value in enumerate(rmse):
        x1 = left + slot * i + slot * 0.22
        x2 = left + slot * (i + 1) - slot * 0.22
        y = bottom - (value / ymax) * (bottom - top)
        draw.rounded_rectangle((x1, y, x2, bottom), radius=8, fill=colors[i])
        draw_centered(draw, ((x1 + x2) / 2, y - 28), f"{value:.2f} C", label_font)
        draw_centered(draw, ((x1 + x2) / 2, bottom + 55), labels[i], label_font)
        draw_centered(draw, ((x1 + x2) / 2, bottom + 125), f"R2  {r2[i]:.3f}", tick_font, "#555555")
    draw_vertical_label(img, (55, (top + bottom) / 2), "Test RMSE  degC", axis_font)
    img.save(path, dpi=(220, 220))


def make_holdout_comparison(path: Path) -> None:
    labels = ["9 point sparse sweep", "49 point dense sweep"]
    values = [42.10800056154824, 0.51842500254158]
    img = Image.new("RGB", (1450, 850), "white")
    draw = ImageDraw.Draw(img)
    title_font = font(ARIAL_BOLD, 40)
    axis_font = font(ARIAL_BOLD, 27)
    tick_font = font(ARIAL, 22)
    label_font = font(ARIAL_BOLD, 26)
    draw_centered(draw, (725, 55), "Independent COMSOL holdout error", title_font)
    left, top, right, bottom = 170, 135, 1320, 690
    draw.line((left, top, left, bottom), fill="#555555", width=3)
    draw.line((left, bottom, right, bottom), fill="#555555", width=3)
    ymax = 45
    for yval in range(0, 46, 5):
        y = bottom - (yval / ymax) * (bottom - top)
        draw.line((left, y, right, y), fill="#E1E1E1", width=2)
        draw.text((left - 18, y), str(yval), font=tick_font, fill="#333333", anchor="rm")
    slot = (right - left) / 2
    colors = ["#B03A2E", "#17365D"]
    for i, value in enumerate(values):
        x1 = left + slot * i + slot * 0.25
        x2 = left + slot * (i + 1) - slot * 0.25
        y = bottom - (value / ymax) * (bottom - top)
        draw.rounded_rectangle((x1, y, x2, bottom), radius=8, fill=colors[i])
        draw_centered(draw, ((x1 + x2) / 2, max(top + 25, y - 30)), f"{value:.3f}%", label_font)
        draw_centered(draw, ((x1 + x2) / 2, bottom + 55), labels[i], label_font)
    draw_vertical_label(img, (60, (top + bottom) / 2), "Relative error  percent", axis_font)
    img.save(path, dpi=(220, 220))


def make_mesh_figure(mesh_path: Path, path: Path) -> None:
    data = json.loads(mesh_path.read_text(encoding="utf-8"))["results"]
    sizes = [float(item["mesh_size"].split("[")[0]) for item in data]
    temps = [float(item["Tmax_K"]) for item in data]
    img = Image.new("RGB", (1450, 850), "white")
    draw = ImageDraw.Draw(img)
    title_font = font(ARIAL_BOLD, 40)
    axis_font = font(ARIAL_BOLD, 27)
    tick_font = font(ARIAL, 22)
    label_font = font(ARIAL_BOLD, 24)
    draw_centered(draw, (725, 55), "Mesh stability of the maximum temperature", title_font)
    left, top, right, bottom = 180, 145, 1300, 690
    draw.line((left, top, left, bottom), fill="#555555", width=3)
    draw.line((left, bottom, right, bottom), fill="#555555", width=3)
    ymin, ymax = 330.1, 330.75
    for idx in range(6):
        val = ymin + idx * (ymax - ymin) / 5
        y = bottom - (val - ymin) / (ymax - ymin) * (bottom - top)
        draw.line((left, y, right, y), fill="#E1E1E1", width=2)
        draw.text((left - 18, y), f"{val:.2f}", font=tick_font, fill="#333333", anchor="rm")
    x_positions = [left + 150, (left + right) / 2, right - 150]
    pts = []
    for x, size, temp in zip(x_positions, sizes, temps):
        y = bottom - (temp - ymin) / (ymax - ymin) * (bottom - top)
        pts.append((x, y))
        draw.text((x, bottom + 45), f"{size:g}", font=label_font, fill="#222222", anchor="mm")
    draw.line(pts, fill="#17365D", width=6)
    for (x, y), temp in zip(pts, temps):
        draw.ellipse((x - 10, y - 10, x + 10, y + 10), fill="#4F81BD", outline="#17365D", width=3)
        draw.text((x, y - 30), f"{temp:.3f} K", font=label_font, fill="#222222", anchor="ms")
    draw_centered(draw, ((left + right) / 2, 790), "Maximum element size  mm", axis_font)
    draw_vertical_label(img, (60, (top + bottom) / 2), "Maximum temperature  K", axis_font)
    img.save(path, dpi=(220, 220))


def set_run_font(run, name: str, size: float, bold: bool = False,
                 italic: bool = False, color: str = TEXT) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 110, start: int = 120,
                     bottom: int = 110, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), BORDER)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    """Keep one logical table row together when Word paginates the document."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_keep_with_next(paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_keep_with_next(p, False)
    r = p.add_run(text)
    set_run_font(r, "Times New Roman", 9.5, italic=True, color="333333")


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.35) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_table(doc: Document, headers: list[str], rows: list[list[str]],
              widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    prevent_row_split(header_row)
    for j, header in enumerate(headers):
        cell = header_row.cells[j]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, "Arial", 9, bold=True, color="FFFFFF")
        if widths:
            cell.width = Inches(widths[j])
    for i, row in enumerate(rows):
        data_row = table.add_row()
        prevent_row_split(data_row)
        cells = data_row.cells
        for j, value in enumerate(row):
            cell = cells[j]
            if i % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(value))
            set_run_font(r, "Times New Roman", 9)
            if widths:
                cell.width = Inches(widths[j])
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)


def add_table_caption(doc: Document, text: str, *, page_break_before: bool = False) -> None:
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    r = p.add_run(text)
    set_run_font(r, "Times New Roman", 9.5, italic=True, color="333333")


def add_math(doc: Document, text: str, number: int | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)
    math_para = OxmlElement("m:oMathPara")
    math_para_pr = OxmlElement("m:oMathParaPr")
    justification = OxmlElement("m:jc")
    justification.set(qn("m:val"), "center")
    math_para_pr.append(justification)
    math_para.append(math_para_pr)
    math_obj = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = text if number is None else f"{text}        ({number})"
    math_run.append(math_text)
    math_obj.append(math_run)
    math_para.append(math_obj)
    p._p.append(math_para)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for element in (fld_begin, instr, fld_sep, text, fld_end):
        run._r.append(element)
    set_run_font(run, "Arial", 9, color="555555")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    set_keep_with_next(p, True)


def suppress_paragraph_borders(paragraph) -> None:
    """Override inherited Word paragraph borders with explicit nil borders."""
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:pBdr"))
    if existing is not None:
        p_pr.remove(existing)
    borders = OxmlElement("w:pBdr")
    for edge_name in ("top", "left", "bottom", "right", "between", "bar"):
        edge = OxmlElement(f"w:{edge_name}")
        edge.set(qn("w:val"), "nil")
        borders.append(edge)
    p_pr.append(borders)


def add_paragraph(doc: Document, text: str, italic: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, "Times New Roman", 10.8, italic=italic)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, "Times New Roman", 10.5)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(14)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 13.5, 12, 6),
        ("Heading 2", 11.5, 9, 4),
        ("Heading 3", 10.8, 7, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption.font.size = Pt(9.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string("333333")

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def build_document() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    benchmark = ROOT / "generated" / "benchmarks" / "busbar_joule_heat"
    dense_csv = benchmark / "builders" / "busbar_dense_training_data.csv"
    training_record = benchmark / "models" / "busbar_tmax_dense.training_record.json"
    mesh_report = benchmark / "mesh_convergence_validation.json"
    workflow_fig = ASSET_DIR / "workflow.png"
    response_fig = ASSET_DIR / "response_heatmap.png"
    comparison_fig = ASSET_DIR / "model_comparison.png"
    holdout_fig = ASSET_DIR / "holdout_comparison.png"
    mesh_fig = ASSET_DIR / "mesh_stability.png"
    make_workflow_figure(workflow_fig)
    make_response_heatmap(dense_csv, response_fig)
    make_model_comparison(training_record, comparison_fig)
    make_holdout_comparison(holdout_fig)
    make_mesh_figure(mesh_report, mesh_fig)

    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("A Traceable Requirement Driven Workflow for Automatic COMSOL Multiphysics Model Construction and Validation")
    suppress_paragraph_borders(title)

    add_heading(doc, "Abstract", 1)
    abstract = (
        "Converting an engineering request into an executable multiphysics model requires decisions about geometry, materials, physics interfaces, boundaries, mesh, study type, solver settings, outputs, and physical acceptance tests. A plausible script or a converged solver is insufficient when these decisions are incomplete or contradictory. "
        "This study presents a traceable requirement driven workflow for automatic COMSOL model construction and validation. It converts natural language requests into versioned modeling contracts, identifies missing or conflicting information, selects physics and study types, generates MATLAB LiveLink and Java artifacts, applies static and runtime gates, and records controlled repairs and physical evidence. "
        "The workflow is implemented in an interactive platform that routes modeling instructions through auditable local stages before using an authenticated read only Codex bridge or a configured external API for advisory assistance. "
        "A fixed 60 request diagnostic benchmark spans solid heat transfer, electric currents, electrothermal coupling, solid mechanics, flow heat coupling, and pressure acoustics. It contains 24 complete, 18 incomplete, 12 parameterized or coupled, and 6 explicitly conflicting requests. The full workflow achieved 100 percent exact physics and coupling selection and no false safe decision on this curated set. A later frozen prospective evaluation used 24 tasks from eight device families excluded by label from development. The dependency compiler achieved 79.17 percent exact interface and coupling selection, compared with 70.83 percent for the existing selector. The version 1.0 gate accepted 37.5 percent of blocked tasks and blocked 25 percent of complete tasks. A separately frozen version 1.1 negation aware layer was then tested on 24 requests from eight further families; it accepted none of 16 blocked tasks, blocked none of eight complete tasks, and matched the exact interface and coupling set for 22 of 24 tasks. A frozen four method development comparison then completed 96 observations on the same project authored requests. The full method made 24 of 24 correct gate decisions; direct and retrieval augmented language model baselines each made 16 of 24 and blocked all eight complete requests, although both matched 23 of 24 physics and coupling labels. These tasks were not independently blinded and did not require COMSOL solves. "
        "Six representative physics templates then completed current workstation MATLAB LiveLink solves with MPH, CSV, preview, execution, and physical validation evidence. In P6, all six verified MPH models were also exported as template specific Java sources, compiled with the COMSOL 6.4 Java toolchain, reconstructed, solved, and saved as fresh MPH artifacts, closing the dual execution path on the same workstation. A three dimensional busbar Joule heating case is retained as an EDA extension and downstream surrogate demonstration. "
        "The EDA extension changed the busbar comparison to a fixed 160 A terminal current, solved an 80 point seven variable Latin hypercube design, and verified five fresh optimization candidates in COMSOL. Three candidates formed a verified mass and power Pareto set, reducing mass by 42.25 to 48.26 percent and Joule loss by 39.49 to 41.21 percent relative to the current driven baseline. "
        "A final sealed P3 evaluation solved 60 independent in domain and 20 out of domain cases. Declared range detection achieved an AUROC of 1.000, but the frozen random forest obtained a maximum temperature RMSE of 31.531 K and failed the surrogate release gate. A subsequent P4 remediation retired the opened P3 final set into training, added physics derived features and 80 adaptively selected COMSOL cases, and evaluated the frozen repair on a newly sealed 60 plus 20 case set. Maximum temperature RMSE decreased to 8.925 K and Joule loss RMSE to 0.338 W, but a symmetric bound released only 2 of 60 in domain cases. P5 then froze a locally scaled one sided conformal upper bound before generating a third 60 plus 20 case final design. The new evaluation achieved a maximum temperature RMSE of 11.807 K, constraint upper coverage of 95 to 100 percent, an out of domain AUROC of 1.000, and zero false safe decisions while releasing 10 of 60 domain cases. The seven variable surrogate therefore passed the declared controlled release gate for the fixed physical scope and input ranges. "
        "The baseline solution satisfied a mesh stability criterion, with a 0.0755 percent change in maximum temperature between the 3 mm and 2 mm maximum element sizes, and an energy balance criterion, with a 0.0473 percent difference between volumetric Joule heating and outward heat loss. "
        "A nine point parameter sweep was sufficient to test the software pipeline but produced a 42.108 percent relative error at an independent holdout condition. "
        "After the sweep was increased to 49 points, a scaled multilayer perceptron achieved an internal test RMSE of 2.315 degrees Celsius and an R squared value of 0.9912. "
        "At the independent COMSOL condition of 25 mV and 7 W per square metre kelvin, the dense surrogate predicted 64.750 degrees Celsius compared with 65.087 degrees Celsius from COMSOL, corresponding to a 0.337 degree Celsius absolute error and a 0.518 percent relative error. "
        "A supplementary two dimensional thermal case used two successive 18 point COMSOL designs: the first exposed a 5.861 K maximum error, and a fresh post retraining holdout reduced the maximum errors to 0.184 K for maximum temperature and 2.744 K for average temperature. These untouched residuals calibrate runtime 95 percent prediction intervals and confidence warnings. "
        "A frozen P7 audit applied one evidence contract to four main EDA cases spanning electrothermal design, electrothermal structural coupling, circuit field coupling, and magnetic frequency response. All four had disjoint training and COMSOL holdout inputs and passed their case specific gates; a fifth electrostatic capacitor case passed as extension evidence. Only the busbar evidence is a fresh current project COMSOL loop, while the remaining cases are identified as preserved independent COMSOL holdouts. "
        "After two EDA decision contracts were frozen, fresh COMSOL runs evaluated two nominal busbar Pareto candidates and a four point MEMS operating grid. Both experiments passed their predeclared gates. The busbar decision selected a 1.872 W, 0.2231 kg design at 160 A, while the MEMS decision selected 5.25 V and 22500 W per square metre kelvin, producing 507.377 K, 1.1197 micrometres displacement, and 15.604 mW. "
        "The results show that explicit contracts, conflict detection, and physics validation prevent unsafe progression that is not visible from code generation or solver completion alone. Surrogate prediction remains an optional downstream capability and is released only for the fixed geometry, materials, boundaries, and input ranges used in validation."
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(abstract)
    set_run_font(r, "Times New Roman", 10.2)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Keywords  ")
    set_run_font(r, "Times New Roman", 10.2, bold=True)
    r = p.add_run("COMSOL Multiphysics; automatic model construction; engineering intent; multiphysics workflow; physical validation; controlled repair; EDA automation; surrogate model")
    set_run_font(r, "Times New Roman", 10.2)

    add_heading(doc, "1 Introduction", 1)
    add_paragraph(doc,
        "Multiphysics finite element models are widely used to evaluate systems in which electrical, thermal, structural, fluid, or acoustic responses interact. The same fidelity that makes these models useful can make repeated evaluation expensive. A parameter study may require tens or hundreds of solutions, while optimization and uncertainty analysis can require many more. Computer experiment methods address this problem by fitting a cheaper response model to deterministic simulation results [1]. Surrogate based analysis extends this idea to engineering design, where the fitted model can screen alternatives or guide additional simulations [2].")
    add_paragraph(doc,
        "A surrogate does not inherit reliability from the finite element solver automatically. Its predictions depend on the quality of the source model, the coverage of the sampled parameter space, the independence of the validation cases, and the consistency between the requested prediction and the model scope. Latin hypercube and related designs were developed to improve the coverage of computer experiments [3], but coverage alone does not verify the physical correctness of the source simulation. Conversely, a numerically converged COMSOL solution does not guarantee that a surrogate fitted to a few selected points will interpolate the response accurately.")
    add_paragraph(doc,
        "Machine learning offers several useful regression families. Ridge regression provides a stable linear reference; random forests can capture nonlinear responses without a prescribed functional form [4]; and multilayer perceptrons can approximate smooth nonlinear mappings when their inputs and outputs are scaled appropriately. Physics informed learning has shown that physical knowledge can improve data driven models [5,6]. The present work takes a more direct engineering approach: the surrogate is trained on solved finite element outputs, while physics based checks are applied to the COMSOL data and to the rules that govern model release. The method does not claim that the neural network solves the governing equations.")
    add_paragraph(doc,
        "This paper addresses the earlier stage: how an engineering intent can be converted into an executable and verifiable COMSOL model without silently inventing decisive inputs. The proposed workflow uses a typed modeling contract, hierarchical physics selection, ten reviewable stages, evidence grades, static checks, physics specific runtime tests, and a bounded repair policy. The principal questions are whether the workflow selects the requested interfaces and study, identifies information that must block execution, prevents contradictory requests from being treated as safe, and leaves artifacts that can be reproduced on another workstation.")
    add_paragraph(doc,
        "Evaluation combines a fixed 60 request, six domain task matrix with six representative COMSOL 6.4 executions. The three dimensional busbar Joule heating model couples current conduction to steady heat transfer, consistent with the Joule Heating interface in COMSOL Multiphysics [8], and serves as the main EDA extension. Its mesh, energy, parameter sweep, independent holdout, and surrogate results show how a generated model can continue into design assessment after the construction and validation gates have passed.")

    add_heading(doc, "1.1 Related work and contribution boundary", 2)
    add_paragraph(doc,
        "Natural language control of simulation and EDA tools already has direct precedents. ChatEDA decomposes natural language EDA requests and generates scripts for an RTL to GDSII flow [9]. MooseAgent generates and iteratively verifies MOOSE input files with a multi agent workflow and a vector database [10]. LayoutCopilot translates analog layout intent into executable commands and design suggestions [11]. ALL FEM combines verified FEniCS code, specialized agents, and runtime feedback across solid, fluid, and multiphysics benchmarks [12]. PDE Agents adds graph augmented retrieval and multi agent orchestration [13], while MOOSEnger combines repository retrieval, executable diagnosis, and persistent memory [14]. FlowFixer EDA applies stage aware log parsing, structured patches, and tool reruns to RTL to GDS failures [15].")
    add_paragraph(doc,
        "Accordingly, this paper does not claim the first language model simulation agent, the first natural language to FEM system, the first knowledge graph augmented simulator, or the first tool feedback repair loop. Its narrower contribution is a COMSOL specific, auditable compiler from requested observables to required interfaces and couplings, combined with explicit conflict, missing information, and validation level stops. An exact local patch is permitted only when it preserves the frozen requirement contract and does not alter geometry, materials, loads, boundary values, physics interfaces, coupling removal, or acceptance thresholds. The prospective evaluation reports failures rather than treating development accuracy as generalization evidence.")

    add_heading(doc, "2 Traceable automatic modeling and surrogate workflow", 1)
    add_heading(doc, "2.1 Workflow and evidence gates", 2)
    add_paragraph(doc,
        "Figure 1 summarizes the proposed automatic modeling workflow. Each stage produces an artifact that can be inspected independently. A typed requirement contract records geometry, materials, boundaries, study, outputs, and units. Missing decisive inputs or explicit contradictions block execution. Hierarchical rules then select physics interfaces and coupling, after which the platform creates a parameterized model plan and mapped MATLAB builder. Static checks precede an approved LiveLink execution. Solver completion is accepted only when case specific physical checks also pass. The resulting evidence package can then support design studies, optimization, or a separately validated surrogate model. Only allowlisted failures receive one controlled repair attempt; other failures remain visible for engineering review.")
    add_figure(doc, workflow_fig,
               "Figure 1  Traceable workflow from engineering requirement to validated COMSOL evidence")

    add_table_caption(doc, "Table 1  Evidence gates used before model execution and downstream release")
    add_table(doc,
              ["Gate", "Required evidence", "Failure response"],
              [
                  ["Finite element model", "Geometry materials boundaries study and outputs are explicitly recorded", "Return to model definition"],
                  ["Numerical verification", "Selected output is stable under mesh refinement", "Refine mesh or solver settings"],
                  ["Physical verification", "Energy conservation or another case specific invariant is satisfied", "Inspect physics selections and expressions"],
                  ["Data integrity", "Finite values unique operating points units and provenance", "Repair or regenerate the sweep"],
                  ["Independent validation", "Fresh COMSOL cases satisfy the declared error threshold", "Collect more evidence or revise the surrogate"],
                  ["Applicability domain", "Requested inputs and fixed model assumptions match the model card", "Refuse prediction and plan new COMSOL cases"],
              ], widths=[1.2, 3.55, 1.85])

    add_heading(doc, "2.2 Data contract and independence", 2)
    add_paragraph(doc,
        "Each sweep row is treated as a numerical experiment. The minimum record contains the input values and units, output expressions and units, model identifier, study type, mesh description, solver status, and source file. The workflow checks missing and nonfinite values, duplicate rows, and overlap between the training and holdout sets. A repeated operating point is not independent evidence even when it appears in a separately named file. The holdout set must be generated by a new COMSOL solve and remain unavailable to the training and candidate selection steps.")
    add_paragraph(doc,
        "The distinction between an internal test split and an independent simulation holdout is central to the method. The internal split compares algorithms under the same data conditions. It can still produce an optimistic estimate because all rows originate from the same sampling design. The independent holdout is a new solved operating point selected within the declared range but excluded from the sweep file. It tests the complete sequence from data preparation through model serialization and prediction.")

    add_heading(doc, "2.3 Candidate models and metrics", 2)
    add_paragraph(doc,
        "Four candidate regressors are compared: a mean predictor, scaled Ridge regression, a random forest, and a multilayer perceptron with scaled inputs and outputs. The mean predictor provides a minimum performance baseline. Ridge regression tests whether a linear response is adequate. The random forest represents a nonlinear tree ensemble, and the multilayer perceptron represents a smooth nonlinear function approximator. Candidate selection uses root mean squared error, mean absolute error, and the coefficient of determination on the internal test split.")
    add_math(doc, "RMSE = √[(1/N) Σᵢ (ŷᵢ − yᵢ)²]", 1)
    add_math(doc, "MAE = (1/N) Σᵢ |ŷᵢ − yᵢ|", 2)
    add_math(doc, "R² = 1 − [Σᵢ (ŷᵢ − yᵢ)² / Σᵢ (yᵢ − ȳ)²]", 3)
    add_paragraph(doc,
        "These statistical metrics are necessary but insufficient. The selected model must also pass the independent COMSOL holdout threshold. Outputs near zero require an absolute error criterion because relative error becomes unstable. The present case has a nonzero temperature output, so both absolute and relative errors are reported.")

    add_heading(doc, "2.4 Applicability domain and controlled prediction", 2)
    add_paragraph(doc,
        "The model card defines the admissible input domain and the fixed modeling assumptions. For the busbar case, the learned mapping is written as")
    add_math(doc, "T̂ₘₐₓ = f(Vₜₒₜ, h)", 4)
    add_paragraph(doc,
        "where the geometry, materials, electrical terminal definitions, convection boundary selection, ambient condition, and stationary study remain fixed. A request is eligible for prediction only when both inputs are present and lie within the validated limits. A change in geometry, material, boundary condition, or coupled physics is treated as a new model scope even if the numerical input values appear similar. Out of range requests produce no surrogate value. Instead, the system records a proposed COMSOL condition for data augmentation, after which a new independent holdout is required.")
    add_paragraph(doc,
        "For an accepted request, the runtime also reports an applicability diagnostic and a prediction interval. The diagnostic normalizes each input to its validated range, records the minimum distance to a range boundary, and, when the training table is available, computes the Euclidean distance to the nearest normalized training condition. The preferred 95 percent interval is the point prediction plus or minus the empirical 95th percentile of absolute residuals from an untouched COMSOL holdout. If point residuals are unavailable, the runtime falls back to a normal approximation based on holdout RMSE or to a reported relative error bound. A deterministic score combines domain position, sample distance, validation status, and interval width. Low confidence or any out of range input sets a COMSOL required flag rather than silently presenting an extrapolation.")

    add_heading(doc, "2.5 Interactive orchestration and assistant boundaries", 2)
    add_paragraph(doc,
        "The workflow is exposed through a local web platform with two user selected operating modes. Offline mode keeps instruction parsing, case memory, staged modeling, surrogate prediction, and artifact inspection on the workstation. API mode provides an advisory language model through either the workstation's authenticated Codex installation or a user configured OpenAI compatible or DeepSeek endpoint. API credentials are stored in a locally ignored directory and configuration responses report only whether a credential is present.")
    add_paragraph(doc,
        "Routing preserves the evidence gates. Requests recognized as automatic modeling tasks enter the local staged workflow before general API dialogue is considered. The user reviews and approves each modeling stage, and local code generation, COMSOL execution, model registration, and validation retain their existing checks. A language model response cannot mark a model as solved or validated. When a local operation fails, the interface may send a redacted error summary and bounded task context to the selected assistant, either after an explicit user action or through an opt in automatic fallback. The returned text is treated as diagnostic guidance and must be applied and verified by the local tools.")
    add_paragraph(doc,
        "The Windows launcher performs a staged cold start. It first reuses or starts the configured COMSOL mphserver, then starts the local execution worker, and finally exposes the web service only after its health endpoint responds. Preview connectivity remains separate from permission to execute new solves. After an installed version compatibility test completed model construction, meshing, a stationary solution, result extraction, and CSV export, external execution was enabled behind the existing model review gate. A single worker lock prevents concurrent local jobs, and a hard timeout terminates the MATLAB process tree if a task stalls. Dependency versions used by the worker are pinned so a future environment rebuild does not silently install an incompatible major release.")

    add_heading(doc, "2.6 Requirement driven multiphysics model construction", 2)
    add_paragraph(doc,
        "The platform also supports a guarded path from a natural language engineering requirement to an executable COMSOL model. The request is first converted into a versioned modeling contract that records the objective, geometry, materials, physics interfaces, boundary and initial conditions, mesh, study, outputs, validation rules, evidence level, and missing information. A ten stage workflow exposes each decision for review. Automatic execution is limited to a verified template with complete critical fields; a new geometry, boundary interpretation, or unsupported interface is assigned a higher review level.")
    add_paragraph(doc,
        "Physics selection uses a fixed catalog of interface names, MATLAB API identifiers, required modules, supported dimensions, study types, material properties, coupling candidates, and validation levels. The selected interfaces and study are returned with matched evidence, confidence, alternatives, and an abstention flag. Before COMSOL starts, a static audit checks parameter definitions, material coverage, empty selections, study creation, result variable prefixes, and output paths. Execution is then recorded in nine stages: server connection, model tree, geometry, selection check, mesh, study, result export, MPH save, and preview generation. Each artifact is fingerprinted with SHA-256. A recoverable lease prevents concurrent MATLAB and COMSOL process trees from using the same local node.")
    add_paragraph(doc,
        "A goal driven physics dependency compiler addresses requests that name desired observables without naming every required COMSOL interface. Temperature, electrical, structural, flow, acoustic, and species goals create explicit dependency edges to candidate interfaces. Joint electrical and temperature goals can require electromagnetic heating; joint temperature and structural goals can require thermal expansion; and fluid temperature goals can require both laminar flow and heat transfer in fluids. Every completed interface records whether it came from the base selector or output dependency completion. The compiler exposes conflicts, unverified couplings, clarification questions, and a blocking decision instead of silently treating a plausible interface list as complete.")
    add_paragraph(doc,
        "Version 1.1 adds an execution critical completeness layer without changing the frozen version 1.0 compiler. It separates clauses that assert available evidence from clauses marked as unavailable, unknown, unspecified, or deferred. The layer checks interface specific boundary fields for thermal, electrical, structural, flow, acoustic, and species models. A missing clause is removed from positive physics evidence, retained in the audit record, mapped to a reason code, and converted into a clarification question. Fixed temperatures on named surfaces count as thermal boundary evidence, and an eigenfrequency structure requires a constraint without an artificial mechanical load requirement.")
    add_paragraph(doc,
        "Physical acceptance is separate from solver completion. The validation layer checks the invariants available for each class, including heat and electric power balance, current continuity, structural reaction and finite response, flow mass balance and pressure drop, acoustic reference frequency and wavelength resolution, and cross field Joule power transfer. Repairs are limited to one diagnosed category per attempt and retain the code and contract difference. Geometry dimensions, material properties, load or boundary values, physics interfaces, coupling removal, and acceptance thresholds require approval. An API assistant cannot bypass these local gates.")

    add_heading(doc, "2.7 Frozen prospective evaluation", 2)
    add_paragraph(doc,
        "After development, version problem1 v1.0.0 froze the dependency compiler, selector, conflict detector, requirement parser, staged workflow, static audit, repair policy, bounded patch implementation, and all task files by SHA 256. The 12 task validation set used four declared device families and the 24 task final set used eight further families; none reused the declared development family labels. The final set was executed once and was not edited after observation. The tasks were authored inside this project after the method freeze and were not independently blinded, so the result is prospective and family disjoint by declaration rather than a third party independent benchmark.")
    add_table_caption(doc, "Table 1A  Frozen validation and final evaluation of the dependency compiler")
    add_table(doc,
              ["Metric", "Validation 12 tasks", "Final 24 tasks"],
              [
                  ["Existing selector exact interfaces and couplings", "6 of 12  50.00%", "17 of 24  70.83%"],
                  ["Dependency compiler exact interfaces and couplings", "8 of 12  66.67%", "19 of 24  79.17%"],
                  ["Dependency compiler exact study", "12 of 12  100%", "23 of 24  95.83%"],
                  ["Full gate false safe rate on blocked tasks", "3 of 8  37.50%", "6 of 16  37.50%"],
                  ["Full gate false block rate on complete tasks", "1 of 4  25.00%", "2 of 8  25.00%"],
              ], widths=[3.15, 1.65, 1.65])
    add_paragraph(doc,
        "The dependency compiler retained an 8.34 percentage point advantage over the existing selector on the final set, providing limited support for cross family transfer. The gate improved on an ungated policy that accepted every blocked task, but the remaining 37.5 percent false safe rate is not adequate for unattended execution. Most false safe cases used negative statements such as an environment or support condition being unavailable; the current keyword evidence did not represent negation. Two complete tasks were conservatively blocked, including an acoustic structure coupling whose implementation remains reference only.")
    add_paragraph(doc,
        "The observed version 1.0 failure was addressed in a separate version 1.1 experiment. Rules were selected on 16 adversarial development requests and then frozen by SHA 256 before any new validation or final result was observed. The validation set contained 12 requests from four further families, and the one time final set contained 24 requests from eight additional families. These tasks were also authored within the project and were not blinded or independently adjudicated. They evaluate requirement interpretation and stopping decisions rather than completed COMSOL builds and solves.")
    add_table_caption(doc, "Table 1B  Frozen version 1.1 negation aware completeness evaluation")
    add_table(doc,
              ["Metric", "Validation 12 tasks", "Final 24 tasks"],
              [
                  ["Version 1.0 false safe rate on blocked tasks", "4 of 8  50.00%", "8 of 16  50.00%"],
                  ["Version 1.1 false safe rate on blocked tasks", "0 of 8  0%", "0 of 16  0%"],
                  ["Version 1.1 false block rate on complete tasks", "0 of 4  0%", "0 of 8  0%"],
                  ["Version 1.1 exact interfaces and couplings", "12 of 12  100%", "22 of 24  91.67%"],
                  ["Version 1.1 exact study", "12 of 12  100%", "24 of 24  100%"],
              ], widths=[3.15, 1.65, 1.65])
    add_paragraph(doc,
        "Version 1.1 made no gate error on either frozen set. The final false safe Wilson 95 percent interval was 0 to 19.36 percent and the false block interval was 0 to 32.44 percent, so zero observed errors must not be read as a zero population error rate. Two blocked final requests missed the exact interface or coupling label, but neither was released for execution. This result closes the specific negative completeness failure under the declared lexicon while leaving unrestricted language, blinded authorship, and end to end solve evidence open.")
    add_paragraph(doc,
        "Six injected static faults tested bounded repair. Every exact patch passed the post repair static audit, preserved the contract hash, and changed on average 17.88 percent of the fixture lines under a maximum three line replacement rule. The comparison represents locality against the scope of replacing a complete artifact. It is not a measured comparison with a frozen external language model regenerating the complete program, so the stronger repair superiority hypothesis remains open.")
    add_table_caption(doc, "Table 1C  Representative frozen contracts executed through COMSOL 6.4")
    add_table(doc,
              ["Evidence stage", "Passed", "Scope"],
              [
                  ["Requirement interfaces couplings and study", "6 of 6", "Six complete frozen contracts"],
                  ["Java compilation solve and new MPH artifact", "6 of 6", "Fresh source workstation rerun"],
                  ["Physics result acceptance", "6 of 6", "Template specific invariants"],
                  ["Full traced pipeline", "6 of 6", "100 percent"],
              ], widths=[2.8, 1.15, 2.85])
    add_paragraph(doc,
        "A separate execution subset linked six complete natural language contracts to the six supported builders. Version 1.1 matched the frozen interfaces, couplings, and study for all six requests. A fresh COMSOL 6.4 Java run compiled every exported builder, reconstructed and solved the first study, and saved a new MPH artifact; all six associated physical result checks passed. SHA 256 values link the contract file, Java report, physical evidence report, and output artifacts. This closes contract to solver traceability for the representative templates, but it does not test arbitrary new geometry synthesis or all 24 version 1.1 final requests.")
    add_paragraph(doc,
        "The repository also freezes a six task complete program regeneration package and an independent blind task intake gate. The regeneration scorer accepts results only when all tasks contain observed model identity, version, operator, run, and time metadata. The blind gate requires at least 24 tasks from eight families, an external author, and a different adjudicator. At the time of reporting, measured external regeneration records were 0 of 6 and accepted external blind tasks were 0 of 24. These assets define acquisition and rejection rules and are not counted as experimental results.")
    add_table_caption(doc, "Table 1D  Frozen four method development comparison on 24 project authored requests")
    add_table(doc,
              ["Method", "Exact physics and coupling", "Exact study", "Correct gate", "False safe and false block"],
              [
                  ["Direct language model", "23 of 24  95.83%", "19 of 24  79.17%", "16 of 24  66.67%", "0 of 16 and 8 of 8"],
                  ["Retrieval augmented language model", "23 of 24  95.83%", "19 of 24  79.17%", "16 of 24  66.67%", "0 of 16 and 8 of 8"],
                  ["Existing selector", "21 of 24  87.50%", "24 of 24  100%", "19 of 24  79.17%", "2 of 16 and 3 of 8"],
                  ["Contract dependency version 1.1", "22 of 24  91.67%", "24 of 24  100%", "24 of 24  100%", "0 of 16 and 0 of 8"],
              ], widths=[1.65, 1.5, 1.0, 1.0, 1.45])
    add_paragraph(doc,
        "The frozen development comparison used the same 24 requests, output schema, and one ephemeral session per language model task. All 48 language model outputs were valid. The direct and retrieval augmented baselines produced the same aggregate counts: they labeled physics and couplings accurately but refused every complete request. The full method made eight more correct gate decisions than either language model baseline; the paired exact McNemar p value was 0.007812 for each comparison. Physics and coupling accuracy did not differ significantly from either language model baseline, with p=1.0. The existing selector made five fewer correct gate decisions than the full method, for p=0.0625 on this sample. Median completion time was 132.2063 s for the direct baseline, 132.8261 s for retrieval, and 0.0001 s for the local full method on the current workstation. Retrieval did not improve any aggregate accuracy count under the frozen corpus. These results support the explicit safety boundary rather than a general language understanding claim.")
    add_paragraph(doc,
        "The associated ablation preserved the same direction. Removing negation aware completeness reduced correct gate decisions from 24 of 24 to 16 of 24 and created eight false safe decisions among 16 blocked requests. Removing the safety gate reduced correct gate decisions to 8 of 24 and created 16 of 16 false safe decisions. The paired gate p values against the full method were 0.007812 and 0.000031, respectively. These four method and ablation observations validate the controlled evaluator and development hypothesis only. The DAC publication gate remains open because the frozen protocol has not yet been repeated on independently authored, sealed P0 tasks.")
    add_paragraph(doc,
        "To prevent the missing independent result from being filled with another project authored set, the P3 protocol is frozen before external task acquisition. It requires 24 requests from eight externally authored and separately adjudicated device families that do not overlap any declared development, validation, or P2 family. Direct and retrieval augmented language model methods each require three independent sessions per task, while the two deterministic methods require one run, for 192 observations. The evaluator is frozen to report task consensus, per family metrics, Wilson 95 percent intervals, exact McNemar tests, paired completion time intervals, repeat agreement, and a failure taxonomy. The protocol, family isolation checker, response contract, and tests are complete, but zero external tasks and zero P3 observations are currently available. No P3 main result is therefore reported.")

    add_heading(doc, "3 Busbar Joule heating case", 1)
    add_heading(doc, "3.1 Geometry materials and boundary conditions", 2)
    add_paragraph(doc,
        "The principal benchmark is a three dimensional busbar assembled from an extruded copper body and titanium alloy bolt domains. The baseline geometry uses a 90 mm length, a 50 mm body width, a 5 mm thickness, a 6 mm circular feature radius, and a nominal 3 mm mesh parameter. One terminal receives an applied electric potential and two terminal faces are grounded. Convective heat flux is applied to the exposed surfaces. Contact resistance, radiative heat transfer, and structural deformation are excluded from the baseline scope.")

    add_table_caption(doc, "Table 2  Baseline busbar model definition")
    add_table(doc,
              ["Category", "Setting", "Value or scope"],
              [
                  ["Geometry", "Busbar body", "Length 90 mm width 50 mm thickness 5 mm"],
                  ["Geometry", "Circular feature and mesh parameter", "Radius 6 mm and nominal mesh parameter 3 mm"],
                  ["Material", "Copper busbar", "Electrical conductivity 5.998e7 S m-1 and thermal conductivity 400 W m-1 K-1"],
                  ["Material", "Titanium beta 21S bolts", "Electrical conductivity 7.407e5 S m-1 and thermal conductivity 7.5 W m-1 K-1"],
                  ["Electrical boundary", "Potential and ground", "Vtot on one terminal and ground on two terminal faces"],
                  ["Thermal boundary", "Convection", "Coefficient h on 25 exposed boundaries; ambient condition inherited from the solved model"],
                  ["Study", "Coupled stationary solve", "Electric Currents Heat Transfer in Solids and electromagnetic heating"],
                  ["Primary output", "Maximum temperature", "Tmax in kelvin and degrees Celsius"],
              ], widths=[1.15, 2.15, 3.3])

    add_heading(doc, "3.2 Governing equations", 2)
    add_paragraph(doc,
        "Current conservation is solved in the conducting domains. With electric field E, current density J, electrical conductivity sigma, and electric potential V, the stationary electrical problem is")
    add_math(doc, "∇ · J = 0", 5)
    add_math(doc, "J = σE     and     E = −∇V", 6)
    add_paragraph(doc,
        "The electromagnetic power dissipation enters the heat equation as a volumetric source")
    add_math(doc, "Qⱼ = J · E", 7)
    add_paragraph(doc,
        "For thermal conductivity k and temperature T, the steady heat balance is")
    add_math(doc, "−∇ · (k∇T) = Qⱼ", 8)
    add_paragraph(doc,
        "Convective heat loss on the selected external surfaces is represented by")
    add_math(doc, "−n · k∇T = h(T − Tₐₘᵦ)", 9)
    add_paragraph(doc,
        "The coupling is nonlinear when material properties depend on temperature. Even with fixed properties, maximum temperature varies nonlinearly with applied voltage because electrical power scales approximately with the square of voltage while convective heat removal depends on the resulting temperature difference.")

    add_heading(doc, "3.3 Mesh and energy verification", 2)
    add_paragraph(doc,
        "The maximum element size was reduced from 4 mm to 3 mm and then to 2 mm. Maximum temperature changed from 330.6341 K to 330.4662 K and 330.2168 K. The relative difference between the last two levels was 0.0755 percent, below the 0.1 percent acceptance criterion used for this benchmark. The Joule power also changed modestly as the electrical solution was refined.")
    add_figure(doc, mesh_fig,
               "Figure 2  Maximum temperature at three mesh resolutions")
    add_table_caption(doc, "Table 3  Mesh convergence results")
    add_table(doc,
              ["Maximum element size", "Maximum temperature K", "Joule power W"],
              [
                  ["4 mm", "330.6341", "3.24655"],
                  ["3 mm", "330.4662", "3.22865"],
                  ["2 mm", "330.2168", "3.20205"],
              ], widths=[2.2, 2.25, 2.15])
    add_paragraph(doc,
        "Energy balance was evaluated on the solved 3 mm model. Volumetric Joule heating was 3.228648 W and outward heat loss was 3.227121 W, giving a relative imbalance of 0.0473 percent. A prior discrepancy was traced to stale geometric boundary identifiers after a rebuild. Replacing the stale list with the actual heat flux feature selection restored the balance. This event illustrates why a smooth output table alone is not sufficient evidence: an invalid boundary selection can generate data that appear numerically regular while representing the wrong physical problem.")

    add_heading(doc, "3.4 Parameter sweep and data sets", 2)
    add_paragraph(doc,
        "Two nested parameter sweeps were used. The initial sweep combined three voltage levels with three heat transfer coefficients for nine operating points. It was intended to test the pipeline and broad response trends. The dense sweep combined seven voltage levels from 10 to 30 mV with seven convection coefficients from 2 to 10 W per square metre kelvin, yielding 49 solved cases. Maximum temperature ranged from 25.626 to 203.812 degrees Celsius. The dense data set contained no missing values or duplicate rows.")
    add_figure(doc, response_fig,
               "Figure 3  Maximum temperature across the dense voltage and convection sweep")
    add_paragraph(doc,
        "The response follows the expected physics. At fixed convection coefficient, increasing voltage increases current, electrical power, and maximum temperature. At fixed voltage, increasing convection reduces maximum temperature. The high voltage and low convection corner has the steepest response, indicating that uniform sampling in raw parameter coordinates may still leave the most nonlinear region relatively underresolved.")

    add_heading(doc, "3.5 Experimental validation protocol", 2)
    add_paragraph(doc,
        "A machine readable protocol was added for future physical temperature validation. Each accepted record must identify the operating condition, measured maximum and average temperatures, the corresponding expanded 95 percent measurement uncertainties, and COMSOL values from the same condition. The validated surrogate is then evaluated with its applicability domain and calibrated prediction interval enabled. For each response, the software reports experimental disagreement with COMSOL and with the surrogate, together with bias, mean absolute error, root mean square error, maximum absolute error, acceptance rate, and overlap between the measurement and surrogate intervals. A row outside the validated surrogate domain cannot pass the release gate.")
    add_paragraph(doc,
        "The current repository contains the schema, blank data template, configurable tolerance and minimum sample policy, command line validator, and web status endpoint. The blank template produces an awaiting measurements state and a failed release flag. No measured values are substituted from simulations, so implementation of the protocol does not constitute experimental evidence.")

    add_heading(doc, "3.6 Fixed current EDA design experiment", 2)
    add_paragraph(doc,
        "The downstream EDA experiment replaced the baseline fixed potential terminal with a fixed current terminal so candidate geometries were compared at the same electrical load. The current driven baseline requested 160 A and returned 160.095 A, a 19.819 mV drop, a maximum temperature of 329.798 K, 3.171 W of Joule loss, 0.4144 kg of copper, and a 0.0473 percent energy imbalance. Nine low, baseline, and high geometry checks confirmed that length, width, thickness, and circular feature radius rebuilt without invalid selections and produced both feasible and infeasible outcomes under the declared constraints.")
    add_paragraph(doc,
        "The design variables were busbar length, width, thickness, circular feature radius, terminal current, convection coefficient, and ambient temperature. An 80 point Latin hypercube design supplied the initial COMSOL data. The nominal optimization fixed current at 160 A, convection at 5 W per square metre kelvin, and ambient temperature at 293.15 K, then minimized copper mass and Joule loss subject to maximum temperature, normalized voltage drop, current density, and geometry constraints. A random forest screened 20,000 candidates. Because none passed the upper 95 percent conservative constraint bound, five mean feasible candidates were treated only as proposals and required fresh COMSOL solutions before acceptance.")
    add_paragraph(doc,
        "Before training, 60 in domain and 20 deliberately out of domain cases were generated with separate seeds and sealed. They were opened only after the model and candidate set were frozen. The final set was evaluated once and was not added to training. Declared variable bounds defined the applicability domain. A candidate was eligible for automatic release only when the upper 95 percent tree ensemble bounds satisfied all physical constraints and the input remained inside every declared range.")
    add_paragraph(doc,
        "P4 treated the opened P3 domain set as diagnostic training evidence and generated a different final design before remediation began. The repair added cross section, surface, resistance, current density, Joule heating, cooling, thermal rise, and geometric ratio features. Four adaptive rounds selected 20 cases each from 12,000 point Latin hypercube pools by combining ensemble disagreement, proximity to the thermal and voltage constraints, and distance from previously selected points. Each response selected between extra trees and random forests with one or two sample leaves using five fold out of fold RMSE. Symmetric 95 percent bounds were frozen from the corresponding out of fold absolute residual quantiles. The new 60 in domain and 20 out of domain final inputs, their hashes, model choices, interval rule, and thresholds were fixed before the final COMSOL runs.")
    add_paragraph(doc,
        "P5 retained the 220 row training set and fitted response functions while changing only the release bound. Because temperature, normalized voltage, and current density are upper limit constraints, two one sided candidates were calibrated from five fold out of fold signed residuals: a global upper residual quantile and a local version scaled by tree ensemble disagreement. The opened P4 final set was permanently retired and used only to select between these predeclared candidates. The local method passed the diagnostic safety gate and provided higher coverage. Its model, calibration report, release thresholds, and SHA-256 values were then frozen before a third, separately seeded 60 in domain plus 20 out of domain final design was generated.")

    doc.add_page_break()
    add_heading(doc, "3.7 Frozen EDA design task contracts", 2)
    add_paragraph(doc,
        "To separate an EDA design problem from a collection of simulation examples, the DAC P1 protocol froze two machine readable task contracts. Each contract records the engineering decision, design and operating variables with units and ranges, objectives, constraints, required outputs, physics interfaces and couplings, validation rules, automatic stop conditions, evidence tier, and source artifact hashes. Definition completion and prospective experiment completion are reported as separate gates.")
    add_table_caption(doc, "Table 3b  Frozen DAC P1 EDA design tasks")
    add_table(doc,
              ["Task", "EDA stage", "Design inputs", "Constraint outputs", "Engineering decision"],
              [
                  ["Busbar electrothermal co-design", "Package and power delivery interconnect", "Length, width, thickness, connection radius; current and cooling", "Temperature, voltage drop, current density, geometry", "Select a mass and loss Pareto geometry or return to COMSOL"],
                  ["MEMS thermal actuator operating design", "MEMS multiphysics co-design", "Drive voltage and surface cooling; fixed lower cooling", "Temperature, displacement, Joule power", "Select a drive and cooling point or reject the operating point"],
              ], widths=[1.25, 1.35, 1.85, 1.55, 1.8])
    add_paragraph(doc,
        "Both task definitions passed the local completeness and traceability checks. A separate experiment manifest then froze the evaluation points, thresholds, selection rules, and required artifacts before execution. The two post-freeze COMSOL experiments completed, and their report timestamps, scenario identifiers, pass states, MPH files, CSV files, MATLAB logs, and SHA-256 values passed machine checks. The research thresholds support controlled comparison and are not product certification limits.")

    add_heading(doc, "4 Results", 1)
    add_heading(doc, "4.1 Candidate model comparison", 2)
    add_paragraph(doc,
        "The 49 cases were divided into 39 training rows and 10 internal test rows. The mean baseline produced a test RMSE of 25.821 degrees Celsius and a negative R squared value. Ridge regression improved the RMSE to 11.455 degrees Celsius, but the residual error showed that a linear mapping was inadequate. The random forest reduced the RMSE to 6.840 degrees Celsius. The scaled multilayer perceptron performed best, with a test RMSE of 2.315 degrees Celsius, an MAE of 1.658 degrees Celsius, and an R squared value of 0.9912. Its RMSE was 91.0 percent lower than that of the mean baseline.")
    add_figure(doc, comparison_fig,
               "Figure 4  Internal test performance of the four candidate regressors")
    add_table_caption(doc, "Table 4  Candidate model performance")
    add_table(doc,
              ["Model", "Test RMSE degC", "Test MAE degC", "Test R2"],
              [
                  ["Mean baseline", "25.821", "23.277", "-0.091"],
                  ["Scaled Ridge", "11.455", "9.980", "0.785"],
                  ["Random Forest", "6.840", "5.426", "0.923"],
                  ["Scaled MLP", "2.315", "1.658", "0.991"],
              ], widths=[2.2, 1.5, 1.5, 1.4])

    add_heading(doc, "4.2 Independent holdout validation", 2)
    add_paragraph(doc,
        "The independent condition used an applied voltage of 25 mV and a convection coefficient of 7 W per square metre kelvin. COMSOL predicted a maximum temperature of 65.087 degrees Celsius. The model trained on the nine point sweep predicted 92.494 degrees Celsius, corresponding to a 27.407 degree Celsius absolute error and a 42.108 percent relative error. The training software operated correctly, but the sparse design did not represent the nonlinear response with sufficient resolution.")
    add_paragraph(doc,
        "The surrogate trained on the 49 point sweep predicted 64.750 degrees Celsius at the same holdout condition. Its absolute error was 0.337 degrees Celsius and its relative error was 0.518 percent. The dense model satisfied the benchmark limits of 3 degrees Celsius absolute error and 5 percent relative error. It was therefore accepted as a fast estimator within the voltage range of 10 to 30 mV and the convection range of 2 to 10 W per square metre kelvin.")
    add_figure(doc, holdout_fig,
               "Figure 5  Independent holdout error before and after increasing parameter coverage")
    add_table_caption(doc, "Table 5  Independent COMSOL holdout comparison")
    add_table(doc,
              ["Training design", "Rows", "Prediction degC", "COMSOL degC", "Absolute error degC", "Relative error"],
              [
                  ["Sparse grid", "9", "92.494", "65.087", "27.407", "42.108%"],
                  ["Dense grid", "49", "64.750", "65.087", "0.337", "0.518%"],
              ], widths=[1.25, 0.55, 1.15, 1.1, 1.35, 1.05])
    add_paragraph(doc,
        "A supplementary implementation case tested whether the same release rule scales beyond one holdout point. An 18 point two dimensional thermal plate design combined centre, near boundary, high temperature, and Latin hypercube conditions. It had zero overlap with the fitting data. The existing model failed the 5 K absolute error gate because the maximum temperature error reached 5.861 K. Those diagnostic points were then transferred to the training data, increasing the fitting table from 51 to 69 rows. They were no longer treated as validation evidence. A second 18 point design generated with a different seed again had zero training overlap and remained untouched during retraining. On this final set, the maximum absolute errors were 0.184 K for maximum temperature and 2.744 K for average temperature. The repaired model therefore passed the declared 5 K gate.")
    add_table_caption(doc, "Table 5b  Expanded independent COMSOL validation for the supplementary thermal model")
    add_table(doc,
              ["Stage", "Independent rows", "Coverage", "Tmax worst error K", "Tavg worst error K", "Decision"],
              [
                  ["Diagnostic holdout", "18", "Centre boundary high temperature and space filling", "5.861", "2.900", "Fail and retrain"],
                  ["Fresh final holdout", "18", "Centre boundary high temperature and space filling", "0.184", "2.744", "Pass"],
              ], widths=[1.25, 0.9, 2.4, 1.0, 1.0, 0.9])

    add_heading(doc, "4.3 Registry and runtime audit", 2)
    add_paragraph(doc,
        "The same release mechanism was applied to additional cases in the local project. The registry contained 29 surrogate records. Twenty seven were validated for their declared scopes, while two earlier models were retained as superseded evidence after failing an expanded validation target. The runtime audit loaded every eligible model, checked the required model card and validation files, compared expected and produced output columns, and repeated the checks after simulated path relocation. All 27 eligible models passed; no artifact or relocation failure was recorded. These results support the portability of the workflow, not the transfer of one fitted model across different physics.")
    add_table_caption(doc, "Table 6  Project level surrogate audit")
    add_table(doc,
              ["Audit measure", "Result"],
              [
                  ["Registered models", "29"],
                  ["Validated models", "27"],
                  ["Superseded historical models", "2"],
                  ["Eligible models passing runtime prediction", "27"],
                  ["Missing artifacts", "0"],
                  ["Relocation failures", "0"],
              ], widths=[4.7, 1.9])

    add_heading(doc, "4.4 Interactive platform verification", 2)
    add_paragraph(doc,
        "The interactive layer was checked separately from the surrogate accuracy experiment. Mode changes were persisted through the local configuration endpoint, and a disconnected service still allowed the interface to return visibly to offline mode while clearly disabling execution dependent actions. The local Codex bridge completed an authenticated live response without a separately entered API key. It ran an ephemeral read only session, while custom OpenAI and DeepSeek adapters were exercised with mocked protocol responses. Credential masking was checked at both the public configuration response and work log boundaries. A cold start on the target Windows workstation reached the web health endpoint, brought the worker online, connected localhost port 2036, and listed 30 MPH artifacts in the observation endpoint.")
    add_paragraph(doc,
        "The installed MATLAB R2023b Update 6 and COMSOL 6.4 environment then completed a full LiveLink compatibility solve. A queue level closure test generated an approved one dimensional heat conduction builder, called MATLAB and COMSOL, solved the model in 47.735 seconds, saved a 1.56 MB MPH file, exported a result CSV, and passed the expected 293.15 to 373.15 K temperature range check. External execution was enabled only after this evidence was recorded. Runtime prediction tests also confirmed that accepted surrogate responses include a domain diagnostic, calibrated intervals, and a confidence grade, while out of range requests are marked low confidence and require COMSOL.")
    add_table_caption(doc, "Table 7  Verification of the interactive orchestration layer")
    add_table(doc,
              ["Component", "Verification", "Observed result"],
              [
                  ["Task routing", "Automatic modeling instruction evaluated before general API dialogue", "Local staged workflow retained priority"],
                  ["Mode control", "API to offline to API transition through the user interface and configuration endpoint", "State changed and persisted"],
                  ["Local Codex bridge", "Authenticated ephemeral read only request without a project API key", "Response returned successfully"],
                  ["Credential handling", "Public configuration and work log inspected after a temporary key save", "No key text was exposed"],
                  ["Cold startup", "COMSOL preview server worker and web health checks executed in sequence", "Port 2036 worker and web service online"],
                  ["Full solve closure", "Approved one dimensional heat task executed through the local queue", "MPH CSV and passed result validation produced in 47.735 s"],
                  ["COMSOL observation", "Connection preview and three dimensional export regression checks", "Focused tests passed"],
                  ["Prediction confidence", "Range distance holdout interval and confidence outputs checked", "Low confidence requests require COMSOL"],
                  ["Regression suite", "Assistant startup server preview and three dimensional preview tests", "All focused checks passed"],
              ], widths=[1.45, 3.25, 1.9])

    add_heading(doc, "4.5 Runtime performance and environment reconstruction", 2)
    add_paragraph(doc,
        "Three queue equivalent MATLAB LiveLink executions on the same 24 logical processor Windows 11 workstation took 18.744, 18.113, and 24.968 s; every run returned 373.15 K. Sampling covered the launched MATLAB tree and persistent COMSOL mphserver tree, with maximum simultaneous working set and private memory of 2803.402 and 3056.199 MB. After warmup, 300 direct surrogate calls had median and 95th percentile latencies of 3.651 and 4.287 ms, while 60 local HTTP calls measured 29.890 and 35.233 ms. Relative to the 18.744 s COMSOL median, the local reference speedups were approximately 5134 and 627 times. These measurements characterize this workstation and fixed case.")
    add_paragraph(doc,
        "Complete sweeps measured the three dimensional busbar with Electric Currents, Heat Transfer in Solids, and Joule Heating. The nine condition verification grid finished in 51.591 s; the full 49 condition training grid finished in 188.039 s. Their median condition solve times were 3.254 and 3.048 s. The dense run reached 3871.012 MB simultaneous working set and 4915.918 MB private memory. All 49 trend and energy checks passed, every input matched the historical training CSV, and the maximum temperature difference was 0.000363 K. These results close the planned multiphysics scale and complete sweep benchmarks for the fixed workstation case.")
    add_table_caption(doc, "Table 8  Local runtime benchmark for the validated thermal plate surrogate")
    add_table(doc,
              ["Path", "Repeated calls", "Median latency", "95th percentile", "Reference speedup"],
              [
                  ["Direct Python", "300", "3.651 ms", "4.287 ms", "5134 x"],
                  ["Local HTTP", "60", "29.890 ms", "35.233 ms", "627 x"],
                  ["COMSOL LiveLink", "3", "18.744 s", "24.345 s", "Reference"],
                  ["3D verification grid", "9 cases", "3.254 s per case", "3.901 s per case", "51.591 s total"],
                  ["3D training grid", "49 cases", "3.048 s per case", "3.425 s per case", "188.039 s total"],
              ], widths=[1.35, 1.1, 1.25, 1.3, 1.35])
    add_paragraph(doc,
        "The Python environment was also reconstructed in a new virtual environment from an exact Windows CPython 3.12 lock file. Thirteen runtime dependencies matched the recorded versions, and the core numerical, machine learning, job serialization, and PDF packages imported successfully. A release manifest records SHA-256 hashes for the paper, generator, validated model, independent validation evidence, surrogate registry, runtime audit, execution configuration, environment lock, and performance report. Local API keys and authentication state are excluded.")

    add_heading(doc, "4.6 Cross physics construction and validation", 2)
    add_paragraph(doc,
        "The requirement driven construction path was evaluated separately from the busbar surrogate experiment. A fixed 35 task intent benchmark covered solid heat transfer, electric currents, solid mechanics, laminar flow, heat transfer in fluids, pressure acoustics, diluted species, four coupling candidates, and stationary, transient, frequency, and eigenfrequency studies. Interface precision, recall, and F1 were 1.000; coupling and study accuracies were 1.000; and no incorrect selection was assigned high confidence. These values measure the declared benchmark and do not establish unrestricted natural language coverage.")
    add_paragraph(doc,
        "An additional 18 request development set tested the new output dependency completion path. A direct keyword baseline matched the complete interface and coupling set for 11 of 18 requests, the prior selector matched 10 of 18, and the dependency compiler matched 18 of 18; study type accuracy was 13 of 18, 18 of 18, and 18 of 18, respectively. The set was authored during method development and is reported only as feasibility evidence. It was not used as an independent final test, and the rules must be frozen before evaluation on new device or structure families.")
    add_paragraph(doc,
        "A second fixed matrix evaluated the complete staged decision path on 60 requests from six release domains. Each domain contained four complete requests, three requests with missing geometry, material, or boundary information, two parameterized or coupled requests, and one explicit conflict. Exact physics and coupling selection and study accuracy were 1.000. Critical missing information recall and precision were both 1.000. All 36 requests eligible for generation reached the mapped MATLAB code generator, while all 24 incomplete or conflicting requests were blocked; the false safe rate was therefore zero on this declared set. The decision benchmark reused six representative template classes and does not claim 60 new finite element solves.")
    add_paragraph(doc,
        "The later frozen version 1.1 experiment targeted the negative wording weakness exposed by version 1.0. On its new 24 request final set, version 1.0 accepted 8 of 16 blocked tasks, whereas version 1.1 accepted none and did not block any of eight complete tasks. Exact interface and coupling matching was 22 of 24 and study matching was 24 of 24. Because the tasks were project authored and no COMSOL solve was required for this gate benchmark, the result supports the bounded completeness method but does not establish independent natural language or end to end construction accuracy.")
    add_table_caption(doc, "Table 9  Deterministic comparison on the fixed 60 request benchmark")
    add_table(doc,
              ["Method", "Exact physics and coupling", "False safe among 24 blocked requests", "Role"],
              [
                  ["Direct keyword proxy", "93.33 percent", "24 of 24", "Ungated deterministic proxy, not a language model result"],
                  ["Rules without case memory", "88.33 percent", "24 of 24", "Physics rules without contract gates"],
                  ["Retrieval plus generation", "100 percent", "6 of 24", "Completeness gate without explicit conflict gate"],
                  ["Full staged workflow", "100 percent", "0 of 24", "Contract completeness conflict audit and validation gates"],
              ], widths=[1.45, 1.45, 1.45, 2.25])
    add_paragraph(doc,
        "The comparison uses deterministic, reproducible baselines. It should not be interpreted as a measured general language model or human expert comparison. In the ablation results, removing the explicit conflict gate left all six contradictory requests false safe. Removing the two static audit fixtures would leave an undefined parameter and an empty selection unblocked. Removing physics validation would treat solver completion for all six representative templates as sufficient without checking temperature bounds, current density, Joule heat, displacement and stress, mass balance and pressure drop, or acoustic reference frequency. Disabling controlled repair left both repair fixtures unresolved.")
    add_paragraph(doc,
        "Six templates were then executed through MATLAB LiveLink on the same Windows workstation with COMSOL 6.4. Every task produced an MPH file, result CSV, preview image, staged execution record, and a passed physics result check. The open tube acoustic model returned 343.0023 Hz against the analytical 343 Hz first mode. The flow heat model passed a 5 percent mass balance criterion with a 1.36 percent relative residual, a positive 0.2401 Pa pressure drop, an outlet temperature within the inlet to wall range, and positive heat transfer to the fluid. The other four templates passed their temperature, voltage and current density, Joule heating, displacement, and stress checks. The flow heat baseline transfers the solved laminar velocity field into Heat Transfer in Fluids with constant properties; it is evidence for this one way baseline rather than for every form of conjugate or temperature dependent flow coupling.")
    add_table_caption(doc, "Table 10  Current workstation evidence for six requirement driven COMSOL templates")
    add_table(doc,
              ["Physics class", "Model and primary check", "Observed result", "Evidence status"],
              [
                  ["Solid heat", "One dimensional steady conduction and temperature bounds", "Maximum temperature 373.15 K", "Passed MPH CSV preview and staged evidence"],
                  ["Electric current", "One dimensional conductor with terminal and ground", "1.0 V maximum and positive current density", "Passed MPH CSV preview and staged evidence"],
                  ["Electrothermal", "One dimensional Joule heating", "294.958 K maximum and positive Joule heat", "Passed MPH CSV preview and staged evidence"],
                  ["Solid mechanics", "Two dimensional loaded rectangle", "22.12 micrometre displacement and 26.26 MPa stress", "Passed MPH CSV preview and staged evidence"],
                  ["Flow heat", "Two dimensional channel mass and heat checks", "1.36 percent mass residual and 0.2401 Pa pressure drop", "Passed one way coupled baseline"],
                  ["Pressure acoustics", "Open tube first eigenfrequency", "343.0023 Hz versus 343 Hz reference", "Passed MPH CSV preview and staged evidence"],
              ], widths=[1.15, 2.2, 1.55, 1.7])
    add_paragraph(doc,
        "A 12 case injected error benchmark achieved 100 percent classification and policy decision accuracy for the fixed cases. Two static repair fixtures, an undefined parameter and an empty selection, failed before repair and passed after one controlled change, giving a 100 percent fixture repair rate with one retry. These tests validate the implemented rules, not autonomous correction of arbitrary COMSOL failures. P6 subsequently exported all six verified MPH models as template specific Java sources. All six sources compiled, reconstructed one component and one study, ran the study through the COMSOL 6.4 Java API, and saved a fresh MPH artifact. The Java execution times ranged from 11.66 to 19.70 seconds on the current workstation. This closes the six template dual path gap without expanding the claim to arbitrary models or another host.")

    add_heading(doc, "4.7 Fixed current design optimization and final release test", 2)
    add_paragraph(doc,
        "All 80 Latin hypercube cases solved successfully. Twenty eight met the declared temperature, normalized voltage, current density, and geometry constraints, while 52 were infeasible. Across this deliberately broad design, maximum temperature ranged from 301.999 to 633.877 K, normalized voltage drop ranged from 7.027 to 54.250 mV, and copper mass ranged from 0.2636 to 0.6746 kg. Every case kept terminal current within one percent of its request and energy imbalance below 0.1 percent.")
    add_paragraph(doc,
        "Five optimization proposals were then solved as fresh COMSOL cases. All five were feasible and improved both baseline objectives; three were nondominated after using the solved rather than predicted values. Relative to the 0.4144 kg and 3.1708 W current driven baseline, the three verified points reduced mass by 48.26, 46.17, and 42.25 percent and reduced Joule loss by 39.49, 40.97, and 41.21 percent. Their maximum temperatures were 318.973, 318.388, and 316.586 K, respectively.")
    add_table_caption(doc, "Table 11  Fresh COMSOL verified Pareto designs at 160 A")
    add_table(doc,
              ["Candidate", "Mass kg", "Mass reduction", "Joule loss W", "Loss reduction", "Maximum temperature K"],
              [
                  ["Pareto 04", "0.2144", "48.26%", "1.9187", "39.49%", "318.973"],
                  ["Pareto 02", "0.2231", "46.17%", "1.8716", "40.97%", "318.388"],
                  ["Pareto 05", "0.2393", "42.25%", "1.8641", "41.21%", "316.586"],
              ], widths=[1.05, 1.0, 1.15, 1.05, 1.15, 1.35])
    add_paragraph(doc,
        "The sealed final evaluation completed 60 in domain and 20 out of domain solutions without execution failure. Range detection achieved an AUROC of 1.000, detected all 20 out of domain cases, and produced no in domain false alarm. The frozen surrogate nevertheless failed the accuracy release gate: maximum temperature RMSE was 31.531 K with R squared 0.220, while Joule loss RMSE was 0.881 W and copper mass RMSE was 0.0308 kg. A mean only gate would have released 25 cases and incorrectly classified three unsafe cases, a 12 percent false safe rate. The implemented upper 95 percent constraint gate released no cases and produced no false safe result. This outcome is safe but has zero surrogate coverage; the seven variable surrogate remains a candidate generator that requires COMSOL verification rather than a released predictor.")
    add_table_caption(doc, "Table 12  Sealed P3 final evaluation")
    add_table(doc,
              ["Measure", "Observed result", "Release decision"],
              [
                  ["COMSOL execution", "80 of 80 solved; current and energy checks passed", "Pass"],
                  ["Out of domain detection", "AUROC 1.000; 100% detection; 0% domain false alarm", "Pass"],
                  ["Maximum temperature", "RMSE 31.531 K; R2 0.220", "Fail 15 K RMSE gate"],
                  ["Joule loss", "RMSE 0.881 W", "Fail 15% relative RMSE gate"],
                  ["Copper mass", "RMSE 0.0308 kg; R2 0.909", "Pass 10% relative RMSE gate"],
                  ["Conservative release", "0 released and 0 false safe", "Safe refusal; no prediction coverage"],
              ], widths=[1.55, 3.25, 1.75])
    add_paragraph(doc,
        "P4 added 80 adaptively selected COMSOL cases in four rounds, increasing the training set from 140 to 220 rows. All added cases passed the terminal current and energy balance checks. The newly sealed final evaluation again solved 60 in domain and 20 out of domain cases. Maximum temperature RMSE decreased from 31.531 to 8.925 K, with an R squared value of 0.963. Joule loss RMSE decreased from 0.881 to 0.338 W, and copper mass RMSE decreased from 0.0308 to 0.00831 kg. Range detection retained an AUROC of 1.000, complete out of domain detection, and zero domain false alarms.")
    add_paragraph(doc,
        "All five reported 95 percent intervals achieved empirical coverage between 96.67 and 100 percent. The temperature interval nevertheless had a mean width of 94.899 K. The conservative constraint rule released two of 60 domain cases and produced no false safe result, corresponding to 3.33 percent useful coverage. This missed the predeclared 10 percent release threshold. A mean prediction rule would have released 28 cases but mislabeled four unsafe cases, a 14.29 percent false safe rate. P4 therefore repaired predictive accuracy without achieving a useful safe release; direct seven variable prediction remains disabled.")
    add_table_caption(doc, "Table 13  Newly sealed P4 remediation evaluation")
    add_table(doc,
              ["Measure", "Observed result", "Release decision"],
              [
                  ["Adaptive COMSOL evidence", "80 of 80 solved; final training set 220", "Pass"],
                  ["Independent COMSOL execution", "60 domain plus 20 out of domain solved", "Pass"],
                  ["Maximum temperature", "RMSE 8.925 K; R2 0.963; PICP 100%", "Pass accuracy and coverage"],
                  ["Joule loss", "RMSE 0.338 W; R2 0.980; PICP 100%", "Pass accuracy and coverage"],
                  ["Out of domain detection", "AUROC 1.000; 100% detection; 0% false alarm", "Pass"],
                  ["Conservative release", "2 of 60 released; 0 false safe; 3.33% coverage", "Fail 10% useful coverage gate"],
              ], widths=[1.55, 3.25, 1.75])
    add_paragraph(doc,
        "P5 evaluated the frozen locally scaled one sided method on the third independent final design. All 80 COMSOL cases solved, met the terminal current and energy balance requirements, and retained the range detector AUROC of 1.000. Maximum temperature RMSE was 11.807 K, Joule loss RMSE was 0.424 W, and copper mass RMSE was 0.00769 kg. The upper bound coverage was 100 percent for maximum temperature, 100 percent for normalized voltage, and 95 percent for current density.")
    add_paragraph(doc,
        "The one sided release rule accepted 10 of 60 domain cases and made no false safe decision, giving 16.67 percent useful coverage and passing the declared 10 percent minimum. A mean only rule would have accepted 21 cases and mislabeled two unsafe cases, a 9.52 percent false safe rate. P5 therefore passed every precommitted execution, accuracy, applicability, upper coverage, false safe, and useful coverage gate. The resulting release remains conditional on the declared seven variable range and fixed physical assumptions.")
    add_table_caption(doc, "Table 14  Third sealed P5 controlled release evaluation")
    add_table(doc,
              ["Measure", "Observed result", "Release decision"],
              [
                  ["COMSOL execution", "60 domain plus 20 out of domain solved", "Pass"],
                  ["Maximum temperature", "RMSE 11.807 K; R2 0.917; upper coverage 100%", "Pass"],
                  ["Voltage and current density", "Upper coverage 100% and 95%", "Pass"],
                  ["Out of domain detection", "AUROC 1.000; 100% detection; 0% false alarm", "Pass"],
                  ["Controlled release", "10 of 60 released; 0 false safe; 16.67% coverage", "Pass 10% useful coverage gate"],
                  ["Mean only diagnostic", "21 released; 2 false safe; 9.52% false safe rate", "Rejected for release"],
              ], widths=[1.55, 3.25, 1.75])

    add_heading(doc, "4.8 Multi case EDA evidence contract", 2)
    add_paragraph(doc,
        "P7 froze a project level evidence contract before auditing the available EDA cases. A main case required local source artifacts with SHA-256 records, the declared input columns, sufficient training and holdout rows, no shared input tuple between training and holdout, and a passed case specific validation report. The project gate required at least four passed main cases from four distinct physics families. This contract checks evidence structure and independence without forcing unlike outputs to share one numerical error threshold.")
    add_paragraph(doc,
        "Four main cases and one extension case passed. The current project busbar case provides the fresh COMSOL optimization and sealed evaluation described above. The MEMS thermal actuator, circuit coupled resistor, and power inductor results are preserved independent COMSOL holdouts from earlier project runs and are labelled as historical evidence rather than new P7 solves. The capacitor result is retained as extension evidence. Training and holdout input overlap was zero for every case.")
    add_table_caption(doc, "Table 15  P7 multi case EDA evidence audit")
    add_table(doc,
              ["Case", "Physics family", "Training and holdout", "Independent result", "Evidence tier"],
              [
                  ["Busbar design", "Electrothermal optimization", "220 and 80; overlap 0", "P5 release passed; 0 false safe", "Fresh current project COMSOL loop"],
                  ["MEMS actuator", "Electrothermal structural", "9 and 4; overlap 0", "Maximum errors 1.41% T, 2.49% displacement", "Preserved COMSOL holdout"],
                  ["FEM resistor", "Circuit field coupling", "5 and 2; overlap 0", "Maximum resistance error 0.37%", "Preserved COMSOL holdout"],
                  ["Power inductor", "Magnetic frequency domain", "13 and 12; overlap 0", "Joint maximum error 1.24%", "Preserved COMSOL holdout"],
                  ["DC capacitor", "Electrostatics extension", "5 and 2; overlap 0", "Maximum capacitance error 0.0023%", "Preserved COMSOL holdout"],
              ], widths=[1.1, 1.45, 1.2, 1.65, 1.35])
    add_paragraph(doc,
        "Passing P7 establishes that the repository can apply one auditable evidence contract across several EDA physics families. It does not establish unrestricted automatic model construction or physical accuracy against experiments. The subsequent DAC P1 contracts convert two cases into explicit EDA decision tasks. Two of two prospective post-freeze COMSOL experiments passed their frozen numerical and artifact gates. The separate P6 publication gate remains open pending real temperature measurements, paired human and external model observations, and destination workstation execution.")

    add_heading(doc, "5 Discussion", 1)
    add_heading(doc, "5.1 Parameter coverage is a release criterion", 2)
    add_paragraph(doc,
        "The contrast between the two holdout results is the clearest finding. A nine point grid reproduced the direction of the response and allowed every software step to complete, yet it missed the independent temperature by more than 27 degrees Celsius. The 49 point grid reduced the error to less than 0.4 degrees Celsius at the same condition. This improvement cannot be attributed only to the selected algorithm because the decisive change was the additional COMSOL evidence. Small data surrogate studies should therefore report the sampling design and the independent error together, rather than presenting a fitted score in isolation.")
    add_paragraph(doc,
        "The earlier 49 point busbar grid resolves only voltage and convection, so its single low error holdout cannot establish a uniform bound across geometry and operating changes. The P3 experiment addressed that evidence gap with seven variables and a sealed 60 point in domain set. Its failure is informative: high cross validation scores for current density, voltage, power, and mass did not guarantee adequate temperature accuracy over the enlarged domain. The supplementary thermal plate experiment shows the correct response to such a result. Diagnostic points may be moved into training, but a different untouched final set is then required. The opened P3 final set cannot be reused as independent evidence for a repaired surrogate.")
    add_paragraph(doc,
        "P4 followed that rule and demonstrated a second limitation of release evidence. Adaptive sampling and physics derived features improved independent temperature accuracy enough to pass the declared RMSE gate, yet accuracy alone did not make the model useful under a conservative constraint policy. The symmetric global residual bound was well calibrated but inefficient near one sided upper constraints. P5 addressed this issue without changing the fitted response functions. It used the retired P4 set only to select a local one sided bound, froze that choice, and obtained independent confirmation on a third final design. This sequence distinguishes method development evidence from final release evidence.")

    add_heading(doc, "5.2 Physics checks protect the training data", 2)
    add_paragraph(doc,
        "A surrogate can reproduce a numerical error as accurately as a physical response. The mesh study limits one source of discretization error, while the power balance tests the consistency of the coupled electrical and thermal fields. The boundary selection incident is especially relevant to automated modeling. Geometry rebuilds can change entity identifiers. If a script continues to integrate heat flux over stale identifiers, the reported power loss can be wrong even when the solver converges. Associating validation expressions with the actual physics feature selection makes the evidence more resilient to geometry updates.")
    add_paragraph(doc,
        "The present workflow uses post hoc physical evidence rather than embedding the governing equations in the loss function. This choice keeps the regression stage simple and allows several model families to be compared. It also means that the fitted function is not guaranteed to satisfy conservation between sampled points. For low dimensional response quantities such as maximum temperature, explicit finite element verification and conservative release limits may be sufficient. Field surrogates or wider parameter spaces would benefit from stronger physical regularization and uncertainty estimation.")

    add_heading(doc, "5.3 Applicability limits and refusal behavior", 2)
    add_paragraph(doc,
        "The validated surrogate should be interpreted as a local response model for one fixed busbar definition. Changing the body dimensions, replacing copper, adding contact resistance, enabling radiation, changing the convection surfaces, or adding thermal stress creates a different mapping. Numerical values produced under those changes may still fall between 25 and 204 degrees Celsius, but they are outside the evidence represented by the model card. Range checks on voltage and convection are therefore necessary but not sufficient; fixed assumptions must also match.")
    add_paragraph(doc,
        "Refusal is treated as a normal engineering outcome. When a request is outside the validated range, the system returns the violated condition and creates a candidate COMSOL sampling plan. A widened model is not released immediately after adding those cases. The augmented training data require a new holdout set that was not used in the update. This rule prevents the evidence used to diagnose a model failure from being reused as if it were an independent confirmation of the repaired model.")

    add_heading(doc, "5.4 Human supervision and assistant boundaries", 2)
    add_paragraph(doc,
        "The conversational assistant is an orchestration aid rather than an evidence source. Its role is to explain requirements, propose corrections, and help interpret a failed local task. Deterministic artifacts from COMSOL, the training pipeline, the independent holdout, and the registry remain authoritative. Keeping modeling requests on the local staged path also prevents a fluent remote answer from bypassing geometry, physics, mesh, or validation checks. This separation is necessary because language model output can be incomplete or incorrect even when it is well written.")
    add_paragraph(doc,
        "The interface therefore exposes the selected mode and provider, requires an explicit opt in for automatic failure handoff, redacts likely credentials, truncates transmitted context, and runs the default local Codex bridge in an ephemeral read only sandbox. These controls reduce accidental disclosure and unintended file changes, but they do not establish the scientific correctness of an assistant response. Any proposed code or modeling change must still be reviewed and executed through the local workflow.")

    add_heading(doc, "5.5 Limitations", 2)
    add_paragraph(doc,
        "Several limitations constrain the conclusions. The released two input busbar surrogate keeps geometry, materials, electrical contacts, ambient conditions, and boundary selections fixed. P5 extends controlled release to seven variables, but only for the declared ranges and the same geometry semantics, materials, contacts, convection roles, and stationary electrothermal formulation. Its one sided bounds describe disagreement with the finite element implementation rather than experimental measurement uncertainty. A three way experimental comparison pipeline and uncertainty aware data contract are implemented, but no measured temperature records are available. Its current report therefore remains awaiting measurements and the work still evaluates agreement with COMSOL rather than physical accuracy. Runtime evidence was collected on one workstation. Resource counters include the MATLAB and mphserver process trees but exclude unrelated operating system processes. Consequently, the reported timing and resource values are local case measurements and not general performance claims.")
    add_paragraph(doc,
        "The assistant integration was tested for routing, connectivity, credential masking, and failure behavior. The frozen P2 experiment additionally measured two language model baselines under a narrow output contract, but it does not establish the scientific accuracy of unrestricted answers. Availability and response time also depend on the selected provider and network. The 35 intent tasks and 60 staged tasks are curated declared sets; the deterministic direct keyword comparison is a proxy and is not a measured unrestricted language model or human expert baseline. The version 1.0 and version 1.1 frozen evaluations and the P2 four method comparison were authored within the project and were not independently blinded. Version 1.1 removed the observed negative completeness gate errors on its 24 task final set, but the zero false safe count has a 19.36 percent Wilson 95 percent upper bound. The separate six contract COMSOL experiment establishes execution traceability only for matured representative builders. The method remains a bounded lexicon and regular expression layer rather than unrestricted language understanding. The bounded repair benchmark verifies static audit recovery and contract preservation on six injected faults, but the frozen external complete program comparison currently has zero measured records. The independent blind task intake also has zero accepted external tasks. The six construction templates are compact baselines rather than evidence for arbitrary industrial geometries. The flow heat template uses one way velocity transfer with constant properties. P6 verified all six Java alternatives on the source workstation, but this does not replace a destination workstation acceptance run. P7 broadens the EDA evidence to four main physics families, but three of those main cases and the capacitor extension reuse preserved independent COMSOL holdouts with small sample counts rather than new P7 solves. Future work should add fresh preregistered solves with larger holdouts for those cases, blinded human and frozen model API comparisons, adversarial requirements, and bidirectional temperature dependent flow coupling. At least five independently recorded surface temperature or thermocouple conditions, with sensor uncertainty and matching COMSOL conditions, should be entered into the implemented protocol before the model can support claims about the physical system rather than the finite element implementation.")
    add_paragraph(doc,
        "The frozen DAC P3 infrastructure does not remove this limitation by itself. It specifies family disjointness, task sealing, repeated language model runs, confidence intervals, paired tests, and failure analysis, but its external task and observation counts remain zero. The manuscript must keep the P2 table labelled as development evidence until the P0 acquisition campaign supplies valid external authorship and adjudication records and all 192 P3 observations are complete.")

    add_heading(doc, "6 Conclusions", 1)
    add_paragraph(doc,
        "This study developed a traceable requirement driven workflow for automatic COMSOL Multiphysics model construction and validation. The workflow separates intent parsing, completeness and conflict checks, physics and study selection, contract approval, code generation, static audit, finite element execution, physics validation, and controlled repair. It records every stage and refuses automatic execution when critical information or contradictory boundaries remain.")
    add_paragraph(doc,
        "The busbar case showed why these stages are needed. The baseline model achieved a 0.0755 percent mesh change in maximum temperature and a 0.0473 percent energy imbalance. A nine point sweep supported the expected response trend but produced a 42.108 percent error at an independent operating point. Expanding the sweep to 49 points and selecting a scaled multilayer perceptron reduced the same holdout error to 0.518 percent. The result supports fast maximum temperature estimation for the fixed benchmark over 10 to 30 mV and 2 to 10 W per square metre kelvin. It does not support extrapolation or changes to geometry, materials, boundaries, or coupled physics.")
    add_paragraph(doc,
        "The central practical conclusion is that a fitted surrogate should be released only with evidence about its source simulation, sampling coverage, independent error, and applicability domain. This requirement is more informative than a high internal R squared value alone and provides a reproducible basis for deciding when the surrogate may answer and when COMSOL must solve a new case.")
    add_paragraph(doc,
        "The interactive implementation preserves this decision boundary by routing modeling operations through local staged tools and using language models only for dialogue and diagnostic assistance. On the fixed 60 request matrix, the full workflow achieved perfect exact physics and coupling selection, study accuracy, and critical missing information recall and precision, generated code for all 36 eligible requests, and produced no false safe decision among 24 blocked requests. Six representative current workstation COMSOL templates completed build, mesh, solve, export, and physics validation. The web interface persists the contract and ten stage state, restores tasks after refresh or service restart, restarts a failed stage, and exports a task archive. A hash manifest release archive and a destination workstation preflight make the implemented scope transferable. The exact Python 3.12 environment was independently rebuilt from a lock file. The busbar and thermal plate experiments further show that the assistant and generator cannot replace mesh evidence, a physical balance, an independent holdout, or the recorded evidence required for release.")
    add_paragraph(doc,
        "The fixed current EDA extension also closed a design optimization loop without confusing candidate screening with evidence. It solved 80 initial designs, verified five new candidates, and retained three COMSOL confirmed Pareto points that reduced both copper mass and Joule loss. The sealed 80 case final evaluation then rejected the seven variable surrogate because its temperature and power errors exceeded the release limits. This negative result demonstrates the intended behavior: validated COMSOL designs remain available, while a surrogate with inadequate independent accuracy is withheld.")
    add_paragraph(doc,
        "The P4 repair added 80 adaptive COMSOL samples and evaluated the frozen model on another sealed 80 case design. Independent temperature, power, and mass errors then passed their declared limits, and every interval coverage and applicability check passed. The model was still withheld because only 3.33 percent of domain cases could be released under the conservative rule. This result separates numerical accuracy, calibration, safety, and practical coverage as distinct release requirements.")
    add_paragraph(doc,
        "P5 completed the release sequence with a frozen locally scaled one sided upper bound and a third independent 80 case COMSOL design. It released 16.67 percent of domain cases with zero false safe decisions and passed all predeclared gates. The seven variable surrogate may therefore support controlled screening inside its declared range. Inputs outside that range or changes to the fixed physical assumptions still require a new COMSOL solution and new validation evidence.")
    add_paragraph(doc,
        "P6 closed the source workstation Java dual path gap: six of six representative templates were exported, compiled, reconstructed, solved, and saved through the COMSOL 6.4 Java API. The repository also freezes a 180 row paired evidence template for human COMSOL, direct model API, and staged platform comparisons, while retaining the existing temperature measurement contract. These acquisition assets do not count as results. The external publication gate remains open until at least five real temperature conditions, all paired comparison records, and a destination workstation acceptance run are observed.")
    add_paragraph(doc,
        "P7 then applied a frozen evidence contract to four main EDA cases and one electrostatic extension. The four main cases covered electrothermal design, electrothermal structural coupling, circuit field coupling, and magnetic frequency response; each had disjoint training and independent COMSOL holdout inputs and passed its declared numerical gate. This broadens the demonstrated repository scope while preserving the distinction between the fresh busbar closed loop and historical holdout evidence.")
    add_paragraph(doc,
        "DAC P1 then froze two explicit EDA decision contracts for busbar power delivery and MEMS thermal actuator operating design. Both definitions passed field, unit, constraint, physics plan, stop condition, artifact, and linked evidence checks. Fresh post-freeze COMSOL runs then passed two of two prospective experiment gates. Two nominal busbar Pareto candidates were solved, and the frozen low-power rule selected the 72.292 mm by 42.251 mm by 3.685 mm design with 318.388 K maximum temperature, 11.698 mV voltage drop, 1.872 W Joule loss, and 0.2231 kg copper mass. Four MEMS points were solved, three were feasible, and the frozen rule selected 5.25 V with 22500 W per square metre kelvin surface cooling, yielding 507.377 K, 1.1197 micrometres displacement, and 15.604 mW. These results cover the declared points and do not establish product qualification or full-domain robustness.")
    add_paragraph(doc,
        "The post freeze requirement experiment gives a more cautious conclusion for automatic model selection. Across 24 tasks from eight new declared device families, the dependency compiler improved exact interface and coupling selection from 70.83 to 79.17 percent. It did not close safe abstention: 37.5 percent of incomplete or conflicting tasks were still accepted and 25 percent of complete tasks were blocked. The method is therefore suitable for continued supervised evaluation, while unattended execution remains disabled for ambiguous or reference only cases.")
    add_paragraph(doc,
        "Version 1.1 then froze a negation aware completeness layer and evaluated another 24 requests from eight new families. It reduced the observed false safe count from 8 of 16 under the version 1.0 gate to 0 of 16 and produced no false block among eight complete requests. Exact interface and coupling matching was 22 of 24. Six further frozen representative contracts completed requirement compilation, COMSOL construction, solution, new MPH artifact creation, and physical acceptance in a fresh source workstation run. This closes same contract execution traceability for the supported templates. Unattended general release still requires independently authored blind tasks and measured results from the frozen external complete program regeneration comparison.")
    add_paragraph(doc,
        "P2 completed the internal four method development comparison with 96 observed rows and 48 valid language model responses. Direct and retrieval augmented baselines each matched 23 of 24 physics and coupling labels but blocked all eight complete requests, yielding 16 of 24 correct gate decisions. The complete contract dependency method matched 22 of 24 labels and made 24 of 24 correct gate decisions, with zero false safe and zero false block observations. This result supports explicit contract and safety gates under the frozen project authored tasks; it does not close the independent publication gate.")
    add_paragraph(doc,
        "DAC P3 now freezes the missing independent evaluation rather than inventing its outcome. Its infrastructure checks external authorship, separate adjudication, eight new case families, task hashes, 192 required repeated observations, case level statistics, confidence intervals, paired tests, repeat agreement, and failure categories. The infrastructure gate passes its automated tests, while the independent main result gate remains open at zero of 24 external tasks and zero of 192 observations.")

    add_heading(doc, "Data and software availability", 1)
    add_paragraph(doc,
        "The repository contains the COMSOL and MATLAB builders, fixed 60 request matrix, benchmark and ablation reports, template and repair evidence, busbar data, model cards, migration preflight, startup scripts, experimental schema, and regression tests. The DAC 2027 P2 protocol, 96 method rows, raw language model audit files, controlled comparison, ablation, and readiness report are under generated/dac2027/p2 with the frozen protocol under configs. DAC 2027 P3 contains the frozen independent evaluation protocol, family isolation and repeated evaluation code, infrastructure manifest, tests, and readiness status; it contains no external result table yet. The earlier project phases P3 to P5 contain the successive busbar optimization, remediation, calibration, frozen designs, and final evaluations. P6 contains six Java sources and execution records, the fixed task comparison template, and separate local and external gate status. P7 contains the frozen multi case configuration, case level hashes, independence checks, validation summaries, and release report. DAC2027 P1 contains the two frozen EDA task contracts, the post-freeze experiment manifest, fresh busbar and MEMS MPH and CSV artifacts, execution logs, reports, and verified hashes. Release manifests record SHA-256 values and exclude local credentials and preferences. A destination workstation must still repeat a live template solve because source artifacts cannot prove its installation compatibility.")

    add_heading(doc, "References", 1)
    refs = [
        "[1] J. Sacks, W. J. Welch, T. J. Mitchell, and H. P. Wynn, Design and analysis of computer experiments, Statistical Science, 4(4), 409-423, 1989. https://doi.org/10.1214/ss/1177012413.",
        "[2] N. V. Queipo, R. T. Haftka, W. Shyy, T. Goel, R. Vaidyanathan, and P. K. Tucker, Surrogate based analysis and optimization, Progress in Aerospace Sciences, 41(1), 1-28, 2005. https://doi.org/10.1016/j.paerosci.2005.02.001.",
        "[3] M. D. McKay, R. J. Beckman, and W. J. Conover, A comparison of three methods for selecting values of input variables in the analysis of output from a computer code, Technometrics, 21(2), 239-245, 1979. https://doi.org/10.1080/00401706.1979.10489755.",
        "[4] L. Breiman, Random forests, Machine Learning, 45, 5-32, 2001. https://doi.org/10.1023/A:1010933404324.",
        "[5] M. Raissi, P. Perdikaris, and G. E. Karniadakis, Physics informed neural networks A deep learning framework for solving forward and inverse problems involving nonlinear partial differential equations, Journal of Computational Physics, 378, 686-707, 2019. https://doi.org/10.1016/j.jcp.2018.10.045.",
        "[6] G. E. Karniadakis, I. G. Kevrekidis, L. Lu, P. Perdikaris, S. Wang, and L. Yang, Physics informed machine learning, Nature Reviews Physics, 3, 422-440, 2021. https://doi.org/10.1038/s42254-021-00314-5.",
        "[7] M. Mitchell, S. Wu, A. Zaldivar, P. Barnes, L. Vasserman, B. Hutchinson, E. Spitzer, I. D. Raji, and T. Gebru, Model cards for model reporting, Proceedings of the Conference on Fairness Accountability and Transparency, 220-229, 2019. https://doi.org/10.1145/3287560.3287596.",
        "[8] COMSOL AB, The Joule Heating Interface, COMSOL Multiphysics Reference Manual, version 6.3, 2024. https://doc.comsol.com/6.3/.",
        "[9] Z. He, H. Wu, X. Zhang, X. Yao, S. Zheng, H. Zheng, and B. Yu, ChatEDA A Large Language Model Powered Autonomous Agent for EDA, arXiv 2308.10204, 2023. https://arxiv.org/abs/2308.10204.",
        "[10] T. Zhang, Z. Liu, Y. Xin, and Y. Jiao, MooseAgent A LLM Based Multi agent Framework for Automating Moose Simulation, arXiv 2504.08621, 2025. https://arxiv.org/abs/2504.08621.",
        "[11] X. Chen et al., LayoutCopilot An LLM Powered Multiagent Collaborative Framework for Interactive Analog Layout Design, IEEE Transactions on Computer Aided Design of Integrated Circuits and Systems, 44(8), 3126-3139, 2025. https://doi.org/10.1109/TCAD.2025.3529805.",
        "[12] R. Deotale, A. Srinivasan, Y. Tian, T. Zhang, P. Vlachos, and H. Gomez, ALL FEM Agentic Large Language models Fine tuned for Finite Element Methods, arXiv 2603.21011, 2026. https://arxiv.org/abs/2603.21011.",
        "[13] S. Adhikari, G. Noorsumar, and O. Jensen, PDE Agents An LLM Orchestrated Multi Agent Framework for Automated Finite Element Simulations with Knowledge Graph Augmented Reasoning, arXiv 2606.07850, 2026. https://arxiv.org/abs/2606.07850.",
        "[14] Z. Abulawi, M. Li, G. Giudicelli, Y. Liu, and C. Permann, Deploying Frontier Agentic Technology in MOOSEnger A Multiphysics Capable AI Assistant, arXiv 2608.15881, 2026. https://arxiv.org/abs/2608.15881.",
        "[15] FlowFixer EDA A Tool Feedback Augmented Large Language Model Agent for Automated Diagnosis and Repair of Open Source RTL to GDS Design Flow Failures, IEEE Access, 14, 140412-140428, 2026. https://doi.org/10.1109/ACCESS.2026.3730896.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 0.9
        r = p.add_run(ref)
        set_run_font(r, "Times New Roman", 8.25)

    properties = doc.core_properties
    properties.title = "A Traceable Requirement Driven Workflow for Automatic COMSOL Multiphysics Model Construction and Validation"
    properties.subject = "English manuscript draft based on the COMSOL small model training project"
    properties.author = ""
    properties.keywords = "COMSOL; surrogate model; Joule heating; independent validation"
    properties.comments = ""

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
